"""Drawing helpers on top of python-pptx. All geometry is in inches."""

from __future__ import annotations

import math
import re
from dataclasses import dataclass

from lxml import etree
from pptx import Presentation
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from . import theme as T

LAYOUT_TITLE_ONLY = 5


# --------------------------------------------------------------------------- #
# presentation scaffolding
# --------------------------------------------------------------------------- #
def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(T.SLIDE_W)
    prs.slide_height = Inches(T.SLIDE_H)
    _apply_theme(prs)
    return prs


def _apply_theme(prs: Presentation) -> None:
    """Write the deck palette and fonts into the theme part."""
    theme_part = prs.slide_masters[0].part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    scheme = root.find(f".//{{{a}}}clrScheme")
    scheme.set("name", "Ontology Review")
    for name, color in T.THEME_COLORS.items():
        slot = scheme.find(f"{{{a}}}{name}")
        for child in list(slot):
            slot.remove(child)
        etree.SubElement(slot, f"{{{a}}}srgbClr", val=T.hex_of(color))
    fonts = root.find(f".//{{{a}}}fontScheme")
    fonts.set("name", "Ontology Review")
    fonts.find(f"{{{a}}}majorFont/{{{a}}}latin").set("typeface", T.HEAD_FONT)
    fonts.find(f"{{{a}}}minorFont/{{{a}}}latin").set("typeface", T.BODY_FONT)
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def add_slide(prs: Presentation, background=None):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_TITLE_ONLY])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = background or T.WHITE
    return slide


def set_title(slide, text, x, y, w, h, size=T.TITLE_PT, color=T.INK, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP):
    """Fill the layout's title placeholder so the outline view stays meaningful."""
    ph = slide.shapes.title
    ph.left, ph.top, ph.width, ph.height = Inches(x), Inches(y), Inches(w), Inches(h)
    tf = ph.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    _zero_margins(tf)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _style_run(run, size=size, color=color, bold=True, font=T.HEAD_FONT)
    return ph


def notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = " ".join(text.split())


# --------------------------------------------------------------------------- #
# text
# --------------------------------------------------------------------------- #
def _zero_margins(tf) -> None:
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0


def _style_run(run, size=T.BODY_PT, color=T.INK, bold=False, italic=False, font=T.BODY_FONT):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = font
    f.color.rgb = color


def _set_bullet(paragraph, color, indent_in=0.25, char="•"):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent_in))))
    pPr.set("indent", str(-int(Inches(indent_in))))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buClr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(buClr, qn("a:srgbClr"), val=T.hex_of(color))
    etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
    etree.SubElement(pPr, qn("a:buChar"), char=char)


def _no_bullet(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    etree.SubElement(pPr, qn("a:buNone"))


Run = tuple  # (text, {style overrides})


def text(slide, x, y, w, h, content, size=T.BODY_PT, color=T.INK, bold=False, italic=False,
         font=T.BODY_FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, bullets=False,
         bullet_color=T.TEAL, space_after=6, line_spacing=None, name=None):
    """Add a text box.

    ``content`` is a string, or a list of paragraphs. A paragraph is a string or
    a list of runs, where a run is a string or ``(text, {"bold": True, ...})``.
    """
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    _zero_margins(tf)
    paragraphs = content if isinstance(content, list) else [content]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        if bullets:
            _set_bullet(p, bullet_color)
        runs = para if isinstance(para, list) else [para]
        for r in runs:
            body, style = (r, {}) if isinstance(r, str) else r
            run = p.add_run()
            run.text = body
            _style_run(
                run,
                size=style.get("size", size),
                color=style.get("color", color),
                bold=style.get("bold", bold),
                italic=style.get("italic", italic),
                font=style.get("font", font),
            )
    return box


# --------------------------------------------------------------------------- #
# shapes
# --------------------------------------------------------------------------- #
def _clean_shape(shape, fill=None, line=None, line_w=1.0, dash=None):
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
        if dash:
            shape.line.dash_style = dash
    shape.shadow.inherit = False
    return shape


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, radius=None, dash=None, name=None,
         shape=None):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE)
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if name:
        s.name = name
    _clean_shape(s, fill, line, line_w, dash)
    s.text_frame.text = ""
    return s


def shape_text(shape, content, size=T.BODY_PT, color=T.INK, bold=False, font=T.BODY_FONT,
               align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.05, italic=False):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    lines = content if isinstance(content, list) else [content]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = line if isinstance(line, list) else [line]
        for r in runs:
            body, style = (r, {}) if isinstance(r, str) else r
            run = p.add_run()
            run.text = body
            _style_run(run, size=style.get("size", size), color=style.get("color", color),
                       bold=style.get("bold", bold), italic=style.get("italic", italic),
                       font=style.get("font", font))
    return shape


def circle(slide, cx, cy, d, fill, label="", size=14, color=T.WHITE, bold=True, font=T.BODY_FONT,
           line=None, name=None):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    _clean_shape(s, fill, line)
    if name:
        s.name = name
    if label:
        shape_text(s, label, size=size, color=color, bold=bold, font=font, margin=0)
    return s


def picture(slide, path, x, y, w=None, h=None, name=None):
    """Raster image; give ``w`` or ``h`` alone to keep the aspect ratio."""
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                   Inches(w) if w else None, Inches(h) if h else None)
    if name:
        pic.name = name
    return pic


def pill(slide, x, y, w, h, label, fill, color=T.WHITE, size=12, bold=True, line=None, font=T.BODY_FONT,
         dash=None):
    s = rect(slide, x, y, w, h, fill=fill, line=line, radius=h / 2, dash=dash)
    shape_text(s, label, size=size, color=color, bold=bold, font=font, margin=0.08)
    return s


def line(slide, x1, y1, x2, y2, color=T.TEAL, width=1.75, dashed=False, arrow=True, name=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dashed:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        ln = c.line._get_or_add_ln()
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    if name:
        c.name = name
    return c


# --------------------------------------------------------------------------- #
# composite components
# --------------------------------------------------------------------------- #
def kicker_title(slide, kicker: str, title: str):
    text(slide, T.MARGIN, 0.42, T.SLIDE_W - 2 * T.MARGIN, 0.3, kicker.upper(), size=T.KICKER_PT,
         color=T.TEAL, bold=True, name="Kicker")
    set_title(slide, title, T.MARGIN, 0.72, T.SLIDE_W - 2 * T.MARGIN, 0.8)


def footer(slide, number: int, dark=False):
    color = T.MINT if dark else T.MUTED
    text(slide, T.MARGIN, 7.02, 8, 0.25, T.DECK_FOOTER, size=10, color=color, name="Footer")
    text(slide, T.SLIDE_W - T.MARGIN - 1, 7.02, 1, 0.25, str(number), size=10, color=color,
         align=PP_ALIGN.RIGHT, name="Slide number")


def card(slide, x, y, w, h, title, body, badge=None, fill=T.MIST, title_color=T.INK,
         body_color=T.INK, badge_fill=T.TEAL, body_size=T.CARD_BODY_PT, title_size=T.CARD_TITLE_PT,
         bullets=False, line_color=None):
    """Tinted rounded card. ``badge`` puts a numbered/glyph circle before the title."""
    rect(slide, x, y, w, h, fill=fill, line=line_color, radius=0.14, name=f"Card: {title}")
    pad = 0.25
    tx = x + pad
    if badge is not None:
        circle(slide, x + pad + 0.22, y + pad + 0.22, 0.44, badge_fill, str(badge), size=14)
        tx = x + pad + 0.6
    text(slide, tx, y + pad + 0.04, x + w - pad - tx, 0.45, title, size=title_size,
         color=title_color, bold=True, anchor=MSO_ANCHOR.MIDDLE if badge is not None else MSO_ANCHOR.TOP)
    body_top = y + pad + (0.62 if badge is not None else 0.5)
    text(slide, x + pad, body_top, w - 2 * pad, y + h - body_top - pad * 0.6, body, size=body_size,
         color=body_color, bullets=bullets, space_after=5)


def takeaway(slide, x, y, w, content, size=15, color=T.INK, dot=T.AMBER):
    """One-line key message with an arrow badge: the deck's recurring 'so what'."""
    circle(slide, x + 0.18, y + 0.18, 0.36, dot, "→", size=14)
    text(slide, x + 0.5, y, w - 0.5, 0.36, content, size=size, color=color, bold=True,
         anchor=MSO_ANCHOR.MIDDLE, name="Takeaway")


def stat(slide, x, y, w, value, label, color=T.TEAL, value_pt=54, label_pt=13, label_color=T.MUTED):
    text(slide, x, y, w, 0.95, value, size=value_pt, color=color, bold=True, font=T.HEAD_FONT)
    text(slide, x, y + 0.95, w, 0.6, label, size=label_pt, color=label_color)


# --------------------------------------------------------------------------- #
# code panel with light syntax highlighting for Turtle / SPARQL
# --------------------------------------------------------------------------- #
_TOKEN = re.compile(
    r"(?P<iri><[^>\s]*>)"
    r"|(?P<string>\"[^\"]*\"(?:@[\w-]+)?)"
    r"|(?P<comment>#.*$)"
    r"|(?P<var>\?\w+)"
    r"|(?P<keyword>@prefix\b|\b(?:PREFIX|SELECT|DISTINCT|WHERE|ORDER BY|CONSTRUCT|ASK|FILTER|NOT EXISTS)\b|(?<![\w:])a(?![\w:]))"
    r"|(?P<pname>[A-Za-z][\w-]*:[\w-]*|:[\w-]+)"
    r"|(?P<number>\b\d+\b)",
    re.MULTILINE,
)

_TOKEN_STYLE = {
    "iri": {"color": T.CODE_IRI},
    "string": {"color": T.CODE_STRING},
    "comment": {"color": T.CODE_COMMENT, "italic": True},
    "var": {"color": T.CODE_VAR, "bold": True},
    "keyword": {"color": T.CODE_KEYWORD, "bold": True},
    "pname": {"color": T.CODE_PNAME},
    "number": {"color": T.CODE_STRING},
}


def _highlight(line_text: str):
    runs, pos = [], 0
    for m in _TOKEN.finditer(line_text):
        if m.start() > pos:
            runs.append((line_text[pos:m.start()], {}))
        runs.append((m.group(0), _TOKEN_STYLE[m.lastgroup]))
        pos = m.end()
    if pos < len(line_text):
        runs.append((line_text[pos:], {}))
    return runs or [(" ", {})]


def code(slide, x, y, w, h, source: str, label: str = "", size=T.CODE_PT, name=None):
    rect(slide, x, y, w, h, fill=T.CODE_BG, radius=0.12, name=name or f"Code: {label}")
    if label:
        text(slide, x + w - 1.6, y + 0.12, 1.4, 0.25, label.upper(), size=10, color=T.CODE_COMMENT,
             bold=True, align=PP_ALIGN.RIGHT)
    box = slide.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.22), Inches(w - 0.5), Inches(h - 0.4))
    tf = box.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    _zero_margins(tf)
    for i, ln in enumerate(source.splitlines()):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.05
        for body, style in _highlight(ln):
            run = p.add_run()
            run.text = body
            _style_run(run, size=size, color=style.get("color", T.CODE_TEXT),
                       bold=style.get("bold", False), italic=style.get("italic", False), font=T.CODE_FONT)
    return box


# --------------------------------------------------------------------------- #
# tables
# --------------------------------------------------------------------------- #
def table(slide, x, y, w, header, rows, col_widths=None, size=13, row_h=0.36, mono_cols=(),
          header_fill=T.TEAL, cell_colors=None, cell_fills=None, name=None):
    """Simple ruled table.

    ``cell_colors`` maps (row, col) -> text colour and ``cell_fills`` maps
    (row, col) -> background; filled cells are also set in bold.
    """
    cell_fills = cell_fills or {}
    n_rows, n_cols = len(rows) + 1, len(header)
    gt = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(row_h * n_rows))
    if name:
        gt.name = name
    tbl = gt.table
    tblPr = tbl._tbl.tblPr
    for attr in ("bandRow", "firstRow"):
        tblPr.set(attr, "0")
    widths = col_widths or [w / n_cols] * n_cols
    for i, cw in enumerate(widths):
        tbl.columns[i].width = Inches(cw)
    for r in range(n_rows):
        tbl.rows[r].height = Inches(row_h)
        for c in range(n_cols):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            is_header = r == 0
            filled = not is_header and (r - 1, c) in cell_fills
            if is_header:
                cell.fill.fore_color.rgb = header_fill
            elif filled:
                cell.fill.fore_color.rgb = cell_fills[(r - 1, c)]
            else:
                cell.fill.fore_color.rgb = T.WHITE if r % 2 else T.MIST
            value = header[c] if is_header else rows[r - 1][c]
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = str(value)
            color = T.WHITE if is_header else (cell_colors or {}).get((r - 1, c), T.INK)
            font = T.CODE_FONT if (c in mono_cols and not is_header) else T.BODY_FONT
            cell_size = size - 1 if font == T.CODE_FONT else size
            _style_run(run, size=cell_size, color=color, bold=is_header or filled, font=font)
    return gt


# --------------------------------------------------------------------------- #
# RDF graph diagrams
# --------------------------------------------------------------------------- #
@dataclass
class Node:
    x: float  # centre
    y: float
    label: str
    kind: str = "resource"  # resource | literal | class | ghost
    w: float = 2.0
    h: float = 0.5


NODE_STYLE = {
    "resource": dict(fill=T.TEAL, color=T.WHITE, line=None),
    "class": dict(fill=T.INK, color=T.WHITE, line=None),
    "literal": dict(fill=T.WHITE, color=T.INK, line=T.TEAL),
    "inferred": dict(fill=T.AMBER_TINT, color=T.INK, line=T.AMBER),
    "warn": dict(fill=T.CLAY_TINT, color=T.INK, line=T.CLAY),
    "var": dict(fill=T.WHITE, color=T.TEAL, line=T.TEAL),     # SPARQL variable
    "ghost": dict(fill=T.WHITE, color=T.MUTED, line=T.MUTED),  # someone else's vocabulary
}


def _boundary(node: Node, dx: float, dy: float) -> tuple[float, float]:
    """Point on the node's rectangle boundary in direction (dx, dy)."""
    hw, hh = node.w / 2, node.h / 2
    tx = hw / abs(dx) if dx else math.inf
    ty = hh / abs(dy) if dy else math.inf
    t = min(tx, ty)
    return node.x + dx * t, node.y + dy * t


def graph(slide, nodes: dict[str, Node], edges, label_size=12, node_size=13, label_bg=T.WHITE):
    """Draw nodes and labelled directed edges.

    ``edges`` items are ``(src, dst, label)`` or ``(src, dst, label, style)``
    where style is "asserted" (teal, solid) or "inferred" (amber, dashed).
    An optional fifth item ``(dx, dy)`` nudges the label off the midpoint.
    """
    for edge in edges:
        src, dst, label = edge[:3]
        style = edge[3] if len(edge) > 3 else "asserted"
        nudge = edge[4] if len(edge) > 4 else (0, 0)
        a, b = nodes[src], nodes[dst]
        dx, dy = b.x - a.x, b.y - a.y
        norm = math.hypot(dx, dy)
        ux, uy = dx / norm, dy / norm
        x1, y1 = _boundary(a, ux, uy)
        x2, y2 = _boundary(b, -ux, -uy)
        inferred = style == "inferred"
        color = T.AMBER if inferred else (T.CLAY if style == "warn" else T.TEAL)
        line(slide, x1, y1, x2, y2, color=color, width=2.0, dashed=inferred,
             name=f"Edge: {label}")
        if label:
            lw = max(0.9, 0.095 * len(label) + 0.25)
            mx, my = (x1 + x2) / 2 + nudge[0], (y1 + y2) / 2 + nudge[1]
            box = rect(slide, mx - lw / 2, my - 0.16, lw, 0.32, fill=label_bg, radius=0.08)
            shape_text(box, label, size=label_size, color=color, bold=True, font=T.CODE_FONT, margin=0.02)
    for key, n in nodes.items():
        st = NODE_STYLE[n.kind]
        is_literal = n.kind == "literal"
        s = rect(slide, n.x - n.w / 2, n.y - n.h / 2, n.w, n.h, fill=st["fill"], line=st["line"],
                 line_w=1.5, radius=None if is_literal else n.h / 2, name=f"Node: {n.label}",
                 dash=MSO_LINE_DASH_STYLE.DASH if n.kind == "inferred" else None)
        shape_text(s, n.label, size=node_size, color=st["color"], bold=not is_literal,
                   font=T.CODE_FONT if not is_literal else T.BODY_FONT, margin=0.04)
