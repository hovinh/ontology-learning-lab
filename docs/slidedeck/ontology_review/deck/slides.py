"""Slide content. One function per slide, in deck order.

Each function takes ``(prs, ex)`` where ``ex`` holds the executed examples from
``examples.run_all()``; code panels and result tables are filled from those, so
the slides always show what actually ran.
"""

from __future__ import annotations

from pathlib import Path

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from . import primitives as P
from . import theme as T
from .primitives import Node

L = T.MARGIN
R = T.SLIDE_W - T.MARGIN
CW = R - L

ASSETS = Path(__file__).resolve().parent.parent / "assets"

PARTS = [
    ("What is an ontology?",
     "A definition, its building blocks, and how it differs from things you already know."),
    ("Why we need ontologies",
     "The problem ontologies solve, and why the Web makes it harder."),
    ("The building blocks, layer by layer",
     "One slide per layer of the W3C stack, then all of them answering one question together."),
    ("Deep dive: RDF, architecture and SPARQL",
     "Triples, global names, merging, and asking questions of a graph."),
    ("Deep dive: meaning through inference",
     "Inferencing, RDF Schema and RDFS-Plus: terms mean what they let you conclude."),
    ("Deep dive: vocabularies in the wild",
     "RDFS-Plus in published vocabularies, and SKOS for thesauri and taxonomies."),
    ("Deep dive: OWL and ontologies at work",
     "Restrictions, sets and counting, and three ontologies used on the Web."),
    ("Deep dive: modeling well",
     "Good and bad practice, OWL 2 profiles, and rules."),
    ("Conclusions",
     "What is actually new, and what to take away."),
]

# The building blocks in Part 3, bottom of the stack first: (layer, role).
LAYERS = [
    ("URI", "Names"), ("RDF", "Statements"), ("RDFS", "Vocabulary"), ("RDFS-Plus", "Sameness"),
    ("OWL", "Definitions"), ("SKOS", "Concepts"), ("Rules", "Derivations"), ("SPARQL", "Questions"),
]


# --------------------------------------------------------------------------- #
# shared slide frames
# --------------------------------------------------------------------------- #
def content(prs, kicker, title):
    slide = P.add_slide(prs)
    P.kicker_title(slide, kicker, title)
    return slide


def glyph(slide, ox, oy, s=1.0, hot=0):
    """Decorative little RDF graph: the deck's motif. ``hot`` picks the inferred edge."""
    pts = [(0.6, 0.9), (2.2, 0.3), (3.6, 1.3), (1.6, 2.3), (3.1, 3.1), (0.5, 3.6), (2.3, 4.2)]
    edges = [(0, 1), (1, 2), (0, 3), (1, 3), (2, 4), (3, 4), (3, 5), (4, 6), (5, 6)]
    for i, (a, b) in enumerate(edges):
        (x1, y1), (x2, y2) = pts[a], pts[b]
        inferred = i == hot % len(edges)
        P.line(slide, ox + x1 * s, oy + y1 * s, ox + x2 * s, oy + y2 * s,
               color=T.AMBER if inferred else T.TEAL, width=2.25 if inferred else 1.5,
               dashed=inferred, arrow=False, name="Motif edge")
    for i, (x, y) in enumerate(pts):
        d = (0.42 if i in (1, 3, 4) else 0.3) * s
        P.circle(slide, ox + x * s, oy + y * s, d, T.MINT if i != 3 else T.AMBER, name="Motif node")


def locator(slide, current):
    """Row of layer pills on the kicker line; the current layer is amber."""
    w, gap = 0.86, 0.05
    x0 = R - len(LAYERS) * w - (len(LAYERS) - 1) * gap
    for i, (name, _) in enumerate(LAYERS):
        x = x0 + i * (w + gap)
        if i == current:
            P.pill(slide, x, 0.36, w, 0.27, name, T.AMBER, color=T.INK, size=9)
        elif i < current:
            P.pill(slide, x, 0.36, w, 0.27, name, T.MINT_SOFT, color=T.TEAL_DARK, size=9)
        else:
            P.pill(slide, x, 0.36, w, 0.27, name, T.WHITE, color=T.MUTED, size=9, line=T.LINE)


def layer_frame(prs, index, title):
    name, role = LAYERS[index]
    slide = content(prs, f"Building block {index + 1} of {len(LAYERS)} · {role} · {name}", title)
    locator(slide, index)
    return slide


def what_it_is(slide, x, y, w, h, body, size=14):
    P.text(slide, x, y, w, 0.3, "WHAT IT IS", size=11, bold=True, color=T.TEAL)
    P.text(slide, x, y + 0.32, w, h - 0.32, body, size=size)


def terms(slide, x, y, w, h, items, heading="KEY TERMS", mono=True, size=13):
    """Bulleted glossary: ``items`` are (term, explanation) pairs."""
    P.text(slide, x, y, w, 0.3, heading, size=11, bold=True, color=T.TEAL)
    style = {"bold": True, "color": T.TEAL_DARK}
    if mono:
        style.update(font=T.CODE_FONT, size=size - 1)
    P.text(slide, x, y + 0.32, w, h - 0.32,
           [[(term, style), ("  " + body, {})] for term, body in items],
           size=size, bullets=True, space_after=5)


def adds_missing(slide, adds, missing, missing_label="Still missing: "):
    """Bottom strip on a layer slide: what the layer contributes, and what it can't do."""
    y, h, w1 = 6.2, 0.62, 5.95
    box = P.rect(slide, L, y, w1, h, fill=T.MINT_SOFT, radius=0.1, name="Adds")
    P.shape_text(box, [[("Adds to the ontology: ", {"bold": True, "color": T.TEAL_DARK}), (adds, {})]],
                 size=13, align=PP_ALIGN.LEFT, margin=0.2)
    box = P.rect(slide, L + w1 + 0.2, y, CW - w1 - 0.2, h, fill=T.AMBER_TINT, radius=0.1, name="Missing")
    P.shape_text(box, [[(missing_label, {"bold": True}), (missing, {})]], size=13, align=PP_ALIGN.LEFT,
                 margin=0.2)


def numbered(slide, x, y, w, h, badge, head, body, fill=T.TEAL, size=13, head_size=16):
    """Numbered definition: badge, bold heading, then body text underneath."""
    P.circle(slide, x + 0.22, y + 0.2, 0.42, fill, str(badge), size=14)
    P.text(slide, x + 0.6, y, w - 0.6, h, [[(head, {"bold": True, "size": head_size})], body], size=size,
           space_after=3)


def progress(slide, current):
    for i in range(len(PARTS)):
        x = L + 0.2 + i * 0.55
        if i < current:
            P.circle(slide, x, 6.35, 0.32, T.TEAL, str(i + 1), size=11)
        elif i == current:
            P.circle(slide, x, 6.35, 0.4, T.AMBER, str(i + 1), size=12, color=T.INK)
        else:
            P.circle(slide, x, 6.35, 0.32, T.INK, str(i + 1), size=11, color=T.MINT, line=T.MINT)


def divider(prs, index):
    title, blurb = PARTS[index]
    slide = P.add_slide(prs, background=T.INK)
    P.text(slide, L, 1.45, 3, 1.4, f"{index + 1:02d}", size=96, color=T.AMBER, bold=True,
           font=T.HEAD_FONT, name="Part number")
    P.text(slide, L, 3.05, 7.4, 0.35, f"PART {index + 1}", size=T.KICKER_PT, color=T.MINT, bold=True)
    P.set_title(slide, title, L, 3.45, 7.6, 1.5, size=40, color=T.WHITE)
    P.text(slide, L, 5.0, 7.2, 0.8, blurb, size=18, color=T.CODE_TEXT)
    progress(slide, index)
    glyph(slide, 8.4, 1.4, 1.05, hot=index * 2 + 1)
    return slide


# --------------------------------------------------------------------------- #
# opening
# --------------------------------------------------------------------------- #
def s_title(prs, ex):
    slide = P.add_slide(prs, background=T.INK)
    P.text(slide, L, 1.75, 7, 0.35, "ONTOLOGY REVIEW", size=T.KICKER_PT, color=T.MINT, bold=True)
    P.set_title(slide, "What is an ontology? From RDF to OWL", L, 2.15, 7.6, 1.9, size=44,
                color=T.WHITE)
    P.text(slide, L, 4.2, 7.2, 1.3,
           "A first introduction: what an ontology is, its building blocks, and how the W3C "
           "standards express them, with runnable examples on one small movie graph",
           size=20, color=T.CODE_TEXT)
    P.pill(slide, L, 6.05, 1.55, 0.38, "asserted", T.TEAL, size=12)
    P.pill(slide, L + 1.7, 6.05, 1.55, 0.38, "inferred", T.AMBER, color=T.INK, size=12)
    P.pill(slide, L + 3.4, 6.05, 1.55, 0.38, "warning", T.CLAY, size=12)
    P.text(slide, L + 5.1, 6.05, 3.0, 0.38, "colour key used throughout", size=12, color=T.MINT,
           anchor=MSO_ANCHOR.MIDDLE)
    glyph(slide, 8.4, 1.3, 1.1, hot=3)
    P.notes(slide, """A review of what an ontology is, built up through the Semantic Web
        standards from RDF to OWL. Every code sample and result in this deck comes from one small
        movie graph, executed with rdflib and owlrl when the deck is built. Colour key: teal is
        what a source asserted, amber is what inference added, and clay marks traps.""")
    return slide


# --------------------------------------------------------------------------- #
# Part 1: what an ontology is
# --------------------------------------------------------------------------- #
def s_definition(prs, ex):
    slide = content(prs, "Part 1 · What is an ontology",
                    "An ontology is a shared, formal model of a domain")
    box = P.rect(slide, L, 1.75, CW, 1.1, fill=T.INK, radius=0.14, name="Definition")
    P.shape_text(box, [
        [("“An ontology is a formal, explicit specification of a shared conceptualization.”",
          {"size": 22, "font": T.HEAD_FONT, "bold": True, "color": T.WHITE})],
        [("The most cited definition: Gruber (1993), refined by Borst (1997) and Studer et al. (1998)",
          {"size": 12, "color": T.MINT})],
    ], align=PP_ALIGN.LEFT, margin=0.35)
    words = [
        ("Conceptualization", "An abstract model of a domain: which kinds of things exist in it and "
                              "how they relate.", "movies, people, and who starred in or directed what"),
        ("Explicit", "The concepts and the rules between them are written down, not left in "
                     "people's heads.", "“every actor is a person” is stated, not assumed"),
        ("Formal", "Written in a language with precise meaning, so software can read it and draw "
                   "conclusions.", "RDF, RDFS and OWL, the W3C standards in this deck"),
        ("Shared", "Agreed by a community, so data from different people and systems is read the "
                   "same way.", "a studio and a fan wiki map onto one model"),
    ]
    w = (CW - 0.9) / 4
    for i, (head, body, eg) in enumerate(words):
        P.card(slide, L + i * (w + 0.3), 3.1, w, 2.3, head,
               [body, [("e.g. " + eg, {"italic": True, "color": T.MUTED})]],
               title_size=17, body_size=13)
    box = P.rect(slide, L, 5.65, CW, 1.05, fill=T.AMBER_TINT, radius=0.12, name="Plain words")
    P.shape_text(box, [[("In plain words: ", {"bold": True}),
                        ("an ontology names the kinds of things in a domain, the relationships between "
                         "them and the rules they follow, so that people and machines interpret the same "
                         "data the same way. Data says what is true of particular things; the ontology "
                         "says what the words in that data mean.", {})]],
                 size=15, align=PP_ALIGN.LEFT, margin=0.3)
    P.notes(slide, """Start with the definition. Gruber's 1993 wording was "an explicit specification
        of a conceptualization"; Borst added formal and shared, and Studer, Benjamins and Fensel
        combined them in 1998. Each word does work: a conceptualization is a simplified view of a
        domain; explicit means written down; formal means a machine can process it; shared means a
        community agrees on it. The word comes from philosophy, the study of what exists.""")
    return slide


def s_anatomy(prs, ex):
    slide = content(prs, "Part 1 · What is an ontology",
                    "The four building blocks of every ontology")
    bw = 6.65
    P.rect(slide, L, 1.75, bw, 2.1, fill=T.MIST, radius=0.12, name="Model band")
    P.text(slide, L + 0.25, 1.85, 4, 0.3, "MODEL  ·  also called the schema or TBox", size=11,
           bold=True, color=T.TEAL)
    P.text(slide, L + 0.25, 2.15, bw - 0.5, 0.5,
           "Axioms here: every :Actor is a :Person; a movie stars actors.", size=12, color=T.MUTED)
    P.rect(slide, L, 4.2, bw, 2.1, fill=None, line=T.LINE, radius=0.12, name="Data band")
    P.text(slide, L + bw - 4.25, 5.95, 4, 0.3, "DATA  ·  also called instances or the ABox", size=11,
           bold=True, color=T.TEAL, align=PP_ALIGN.RIGHT)
    nodes = {
        "movie": Node(1.75, 3.2, ":Movie", kind="class", w=1.2),
        "actor": Node(4.15, 3.2, ":Actor", kind="class", w=1.15),
        "person": Node(6.55, 3.2, ":Person", kind="class", w=1.2),
        "giant": Node(1.75, 4.95, ":Giant", w=1.1),
        "dean": Node(4.15, 4.95, ":JamesDean", w=1.7),
        "year": Node(1.75, 5.85, "1956", kind="literal", w=0.8, h=0.4),
    }
    P.graph(slide, nodes, [
        ("movie", "actor", ":starring"),
        ("actor", "person", "subClassOf"),
        ("giant", "dean", ":starring"),
        ("giant", "movie", "rdf:type", "asserted", (0, 0.05)),
        ("dean", "actor", "rdf:type", "inferred", (0, 0.05)),
        ("giant", "year", ":releaseYear", "asserted", (0.75, 0)),
    ], label_size=11)
    for n, (x, y) in enumerate([(0.9, 3.2), (0.9, 4.95), (2.95, 4.58), (5.4, 2.82)], start=1):
        P.circle(slide, x, y, 0.34, T.AMBER, str(n), size=12, color=T.INK)
    P.text(slide, 5.15, 4.75, 2.0, 0.5, "dashed amber = inferred, not typed in", size=11, italic=True,
           color=T.AMBER)
    x0, w = 7.65, R - 7.65
    blocks = [
        ("Classes", "Kinds of things, treated as sets of members: :Movie, :Person, :Actor. Classes form "
                    "hierarchies: every :Actor is also a :Person."),
        ("Individuals", "The particular things, also called instances: the film :Giant, the actor "
                        ":JamesDean. rdf:type says which classes an individual belongs to."),
        ("Properties", "Relationships between individuals (:Giant :starring :JamesDean) and attributes "
                       "with plain values (:releaseYear 1956). A property is not owned by one class."),
        ("Axioms", "Statements a machine can reason with: subclasses, domains and ranges, definitions, "
                   "constraints. :JamesDean is inferred to be an :Actor because he stars in a movie."),
    ]
    for i, (head, body) in enumerate(blocks):
        numbered(slide, x0, 1.75 + i * 1.15, w, 1.1, i + 1, head, body, fill=T.AMBER)
    P.takeaway(slide, L, 6.45, CW,
               "Classes, properties and axioms make the model; add individuals and you have a knowledge graph.")
    P.notes(slide, """Every ontology is built from the same four blocks. Classes are kinds of things,
        understood as sets. Individuals are the things themselves. Properties relate individuals to
        each other or to plain values, and exist independently of classes. Axioms are the statements
        about classes and properties that let a reasoner infer new facts. The top band is the model,
        often called the TBox (terminology); the bottom band is the data, the ABox (assertions).
        Some authors call only the model the ontology. Everything is named with URIs and written as
        RDF triples, so model and data live in one graph.""")
    return slide


def s_spectrum(prs, ex):
    slide = content(prs, "Part 1 · What is an ontology",
                    "From word lists to ontologies: each step adds meaning")
    P.text(slide, L, 1.8, 4.8, 1.2,
           "You have probably used the lower steps already. Each adds structure to the one before; "
           "an ontology adds meaning a machine can reason with.", size=15, italic=True, color=T.MUTED)
    steps = [
        ("Controlled vocabulary", "An agreed list of terms, so everyone uses the same words.",
         "Drama, Western, Comedy", T.MIST, T.INK),
        ("Taxonomy", "+ a hierarchy of broader and narrower terms.",
         "Teen drama → Drama", T.MINT_SOFT, T.INK),
        ("Thesaurus", "+ synonyms and related terms.",
         "Teen drama, also “Coming-of-age”; related to Family drama", T.MINT, T.INK),
        ("Data model or schema", "+ typed records and fields, built for one application.",
         "table Movie(title, year, director)", T.TEAL, T.WHITE),
        ("Ontology", "+ formal meaning a machine can reason with, global names, and reuse across "
                     "applications.", "every Actor is a Person; one registry id means one person",
         T.INK, T.WHITE),
    ]
    w, gap, base = (CW - 4 * 0.15) / 5, 0.15, 6.1
    for i, (name, adds, eg, fill, color) in enumerate(steps):
        h = 2.75 + i * 0.42
        x = L + i * (w + gap)
        P.rect(slide, x, base - h, w, h, fill=fill, radius=0.12, name=f"Step: {name}")
        P.text(slide, x + 0.2, base - h + 0.18, w - 0.4, 0.7, name, size=17, bold=True, color=color,
               font=T.HEAD_FONT)
        P.text(slide, x + 0.2, base - h + 0.95, w - 0.4, 1.0, adds, size=13, color=color)
        P.text(slide, x + 0.2, base - 0.95, w - 0.4, 0.85, eg, size=12, italic=True,
               color=T.AMBER if fill in (T.INK, T.TEAL) else T.TEAL_DARK)
    P.takeaway(slide, L, 6.38, CW,
               "A knowledge graph is an ontology plus the data described with it, stored as one graph.")
    P.notes(slide, """Placing ontologies among familiar artifacts. A controlled vocabulary fixes the
        words. A taxonomy arranges them in a hierarchy. A thesaurus adds synonyms and related terms;
        SKOS, building block 6, publishes these three on the Web. A data model or database schema
        types records for one application. An ontology adds formal meaning that a reasoner can use,
        global names, and the expectation that other applications will reuse and extend it.""")
    return slide


def s_vs_database(prs, ex):
    slide = content(prs, "Part 1 · What is an ontology", "An ontology is not a database schema")
    P.text(slide, L, 1.7, CW, 0.45,
           "Both describe data, but they assume different worlds. Knowing the difference avoids most "
           "beginner surprises.", size=16, color=T.MUTED)
    rows = [
        ("Built for", "Storing and validating data for one application",
         "Sharing meaning across applications and sources"),
        ("A missing fact", "Is false: the closed-world assumption", "Is unknown: the open-world assumption"),
        ("Names", "Local table and column names", "Global URIs that anyone can reuse"),
        ("Data that breaks a rule", "Rejected with a constraint error",
         "Used to infer something; validation is a separate step (SHACL)"),
        ("One thing, two ids", "Two rows until someone deduplicates",
         "Declared or inferred to be the same with owl:sameAs"),
        ("Changing the model", "A migration; the schema comes before the data",
         "Add triples at any time; model and data are both triples"),
        ("Answers come from", "Only what was inserted", "What was asserted plus what can be inferred"),
    ]
    P.table(slide, L, 2.25, CW, ("", "Database schema", "Ontology"), rows,
            col_widths=[2.5, 4.4, CW - 6.9], size=13, row_h=0.46,
            cell_fills={(r, 0): T.MIST for r in range(len(rows))})
    P.takeaway(slide, L, 6.4, CW,
               "They work together: R2RML maps tables to RDF, and OWL 2 QL is designed to run over databases.")
    P.notes(slide, """Most people arrive knowing relational databases, so the contrast helps. A
        schema constrains data for one application and treats anything absent as false. An ontology
        describes meaning for many applications, treats absent facts as unknown, and infers rather
        than rejects. Validation of RDF data is done with SHACL, a separate W3C standard. Unlike
        object-oriented classes, ontology classes are sets: one individual can be in many classes,
        and properties exist on their own.""")
    return slide


def s_why(prs, ex):
    slide = content(prs, "Part 2 · Why we need ontologies", "The problem: same world, different words")
    lw = 6.0
    P.text(slide, L, 1.72, lw, 0.3, "SOURCE A  ·  studio catalogue", size=11, bold=True, color=T.TEAL)
    P.table(slide, L, 2.02, lw, ("movie", "starring"),
            [("Giant", ":JamesDean"), ("East of Eden", ":JamesDean")],
            col_widths=[2.6, lw - 2.6], size=13, row_h=0.36, mono_cols=(1,))
    P.text(slide, L, 3.27, lw, 0.3, "SOURCE B  ·  fan wiki", size=11, bold=True, color=T.TEAL)
    P.table(slide, L, 3.57, lw, ("page", "type", "bornIn"),
            [("fan:Dean_James", "fan:Star", "Marion, Indiana")],
            col_widths=[2.2, 1.5, lw - 3.7], size=13, row_h=0.36, mono_cols=(0, 1))
    box = P.rect(slide, L, 4.5, lw, 2.2, fill=T.CLAY_TINT, radius=0.12, name="Question")
    P.shape_text(box, [
        [("“Where was the star of Giant born?”", {"bold": True, "size": 16})],
        [("A person joins these at a glance. A program must be told three things:", {})],
        [("1   ", {"bold": True, "color": T.CLAY}), ("Identity: fan:Dean_James and :JamesDean are the "
                                                    "same person", {})],
        [("2   ", {"bold": True, "color": T.CLAY}), ("Vocabulary: a fan:Star is what the studio calls an "
                                                    ":Actor", {})],
        [("3   ", {"bold": True, "color": T.CLAY}), ("Meaning: every :Actor is a :Person", {})],
        [("An ontology states these once, in a form software applies to every record.",
          {"italic": True})],
    ], size=13, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=0.25)
    x0, w = 7.0, R - 7.0
    P.text(slide, x0, 1.72, w, 0.3, "WHAT AN ONTOLOGY GIVES YOU", size=11, bold=True, color=T.TEAL)
    gains = [
        ("Shared vocabulary", "Terms with one agreed meaning, written down, across teams and systems."),
        ("Integration", "Map other sources' terms onto yours instead of rewriting their data."),
        ("Inference", "Derive facts nobody typed in: James Dean is a Person."),
        ("Consistency checks", "A reasoner flags statements that contradict the model."),
        ("Reuse", "Publish the model; others extend it instead of starting over."),
    ]
    for i, (head, body) in enumerate(gains):
        numbered(slide, x0, 2.1 + i * 0.74, w, 0.72, i + 1, head, body, size=13, head_size=14)
    box = P.rect(slide, x0, 5.85, w, 0.85, fill=T.AMBER_TINT, radius=0.12, name="Where")
    P.shape_text(box, [[("Where you meet them: ", {"bold": True}),
                        ("schema.org markup read by search engines; Gene Ontology and SNOMED CT in "
                         "biology and medicine; enterprise knowledge graphs; context for LLM apps.", {})]],
                 size=12, align=PP_ALIGN.LEFT, margin=0.2)
    P.notes(slide, """Motivation, after the definition. Two sources describe the same actor with
        different identifiers and different words. A person joins them without noticing; a program
        needs to be told about identity, vocabulary and meaning. An ontology writes those down once,
        so every record benefits. The same pattern scales from two files to whole industries:
        schema.org for the Web, the Gene Ontology and SNOMED CT in the life sciences.""")
    return slide


def s_stack(prs, ex):
    slide = content(prs, "Part 3 · The building blocks",
                    "Ontology languages come in layers")
    P.text(slide, L, 1.75, 7.5, 0.7,
           "Each layer adds terms on top of the one below, and a model can stop at any layer: "
           "pick how much meaning the problem needs.", size=15, italic=True, color=T.MUTED)
    P.line(slide, 1.0, 4.35, 9.15, 2.3, color=T.AMBER, width=2.5, dashed=True)
    layers = [
        ("RDF", "Blocks 1–2", "Global names (URIs) and triples; merging is a union", T.MIST, T.INK, 1.8),
        ("RDFS", "Block 3", "Classes, subclasses, property hierarchies, domain and range", T.MINT_SOFT,
         T.INK, 2.4),
        ("RDFS-Plus", "Block 4", "Sameness, inverse and functional properties", T.TEAL, T.WHITE, 3.0),
        ("OWL", "Block 5", "Restrictions, set operations, cardinality, profiles", T.INK, T.WHITE, 3.6),
    ]
    w, gap, base = 2.35, 0.15, 6.55
    for i, (name, chs, desc, fill, color, h) in enumerate(layers):
        x = L + i * (w + gap)
        P.rect(slide, x, base - h, w, h, fill=fill, radius=0.12, name=f"Layer: {name}")
        P.text(slide, x + 0.2, base - h + 0.18, w - 0.4, 0.4, name, size=20, bold=True, color=color,
               font=T.HEAD_FONT)
        P.text(slide, x + 0.2, base - h + 0.6, w - 0.4, 0.3, chs, size=12, bold=True,
               color=T.AMBER if fill in (T.INK, T.TEAL) else T.TEAL)
        P.text(slide, x + 0.2, base - h + 0.92, w - 0.4, 1.3, desc, size=14, color=color)
    sx = L + 4 * (w + gap) + 0.15
    side = [
        ("SKOS", "Block 6: concept schemes, built on RDFS-Plus"),
        ("Rules", "Block 7: if-then logic beyond OWL"),
        ("SPARQL", "Block 8: queries every layer"),
    ]
    for i, (head, body) in enumerate(side):
        P.card(slide, sx, 1.85 + i * 1.6, R - sx, 1.45, head, body, title_size=17, body_size=14)
    P.notes(slide, """Ontology modeling on the Web is a family of layered languages that differ
        in expressivity. RDF makes and merges simple statements. RDFS adds classes, subclasses
        and properties. RDFS-Plus is a practical subset of OWL, enough for most integration work.
        OWL brings the expressivity of logic. Beside the ladder, SKOS organizes vocabularies, rules
        cover if-then logic, and SPARQL queries all of it. The next eight slides take one building
        block each.""")
    return slide


# Pixel boxes in assets/semantic_web_stack.png (600 x 630) of the layers this review covers,
# labelled with their building-block number, and of the layers it only sketches, lettered.
W3C_LAYERS = [
    ("URI/IRI", (12, 567, 482, 603), "1"),
    ("Data interchange: RDF", (12, 440, 253, 544), "2"),
    ("RDFS", (167, 367, 345, 416), "3"),
    ("Ontology: OWL", (167, 285, 345, 346), "5"),
    ("Rule: RIF", (368, 285, 481, 416), "7"),
    ("Query: SPARQL", (31, 285, 144, 416), "8"),
]
W3C_UNCOVERED = [
    ("User interface & applications", (12, 12, 573, 62), "A"),
    ("Trust", (322, 84, 573, 134), "B"),
    ("Proof", (268, 157, 481, 226), "C"),
    ("Unifying logic", (103, 230, 399, 261), "D"),
    ("Crypto", (504, 157, 573, 603), "E"),
    ("XML", (276, 504, 481, 544), "F"),
]


def stack_picture(slide, boxes, color, label_color):
    """The W3C stack image with outlined, badged boxes over the given layers."""
    ix, iy, ih = L, 1.75, 4.8
    k = ih / 630
    P.picture(slide, ASSETS / "semantic_web_stack.png", ix, iy, h=ih, name="W3C Semantic Web Stack")
    for layer, (x1, y1, x2, y2), badge in boxes:
        x, y, w, h = ix + x1 * k, iy + y1 * k, (x2 - x1) * k, (y2 - y1) * k
        P.rect(slide, x - 0.03, y - 0.03, w + 0.06, h + 0.06, line=color, line_w=2.5, radius=0.05,
               name=f"Layer: {layer}")
        P.circle(slide, x, y, 0.3, color, badge, size=12, color=label_color, name=f"Badge {badge}")
    P.text(slide, ix, iy + ih + 0.06, 600 * k, 0.25,
           "Semantic Web Stack: W3C, public domain, via Wikimedia Commons", size=10, color=T.MUTED)
    return ix, iy, k


def s_stack_w3c(prs, ex):
    slide = content(prs, "Part 3 · The building blocks", "The same layers in the W3C Semantic Web Stack")
    ix, iy, k = stack_picture(slide, W3C_LAYERS, T.AMBER, T.INK)
    P.pill(slide, ix + 256 * k - 0.55, iy + 356 * k - 0.11, 1.1, 0.22, "4 · RDFS-Plus", T.TEAL,
           size=10, line=T.WHITE)
    x0 = 5.75
    P.text(slide, x0, 1.75, R - x0, 0.4, "Numbers are the building blocks on the next slides",
           size=15, bold=True)
    rows = [
        ("1", "URI/IRI", "Global names", "4"),
        ("2", "RDF", "Triples, merging", "4"),
        ("3", "RDFS", "Classes, hierarchies", "5"),
        ("4", "RDFS-Plus (a slice of OWL)", "Sameness", "5"),
        ("5", "Ontology: OWL", "Definitions, constraints", "7"),
        ("6", "SKOS (on RDFS, not drawn)", "Concept schemes", "6"),
        ("7", "Rule: RIF", "If-then rules", "8"),
        ("8", "Query: SPARQL", "Questions", "4"),
    ]
    P.table(slide, x0, 2.2, R - x0, ("Block", "Layer", "Adds", "Deep dive"), rows,
            col_widths=[0.75, 2.75, 2.2, R - x0 - 5.7], size=13, row_h=0.43,
            cell_fills={(r, 0): T.AMBER_TINT for r in range(len(rows))})
    P.takeaway(slide, x0, 6.2, R - x0,
               "An ontology lives mid-stack: RDFS and OWL, on top of RDF.", size=14)
    P.notes(slide, """The W3C Semantic Web Stack, the layer cake, places the same languages in the
        wider Web architecture. Amber outlines mark the layers this review covers, numbered by
        building block; the last column gives the deep-dive part. RDFS-Plus is not a W3C layer: it is
        a practical subset of OWL between RDFS and OWL. SKOS is a vocabulary built on RDFS and is
        not drawn in the diagram. The next slide sketches the layers we do not cover.""")
    return slide


def s_beyond_stack(prs, ex):
    slide = content(prs, "Part 3 · The building blocks", "The rest of the stack, in brief")
    stack_picture(slide, W3C_UNCOVERED, T.CLAY, T.WHITE)
    x0 = 5.75
    P.text(slide, x0, 1.72, R - x0, 0.4,
           "Not covered in this review, but useful to recognize", size=15, bold=True)
    layers = [
        ("User interface & applications", "What people actually use: search, dashboards, chat "
         "assistants. Built on the layers below; no standard of its own."),
        ("Trust", "Deciding which sources and conclusions to believe when anyone can say anything. "
         "Draws on provenance (W3C PROV), proofs and signatures; still largely open."),
        ("Proof", "A record of how a conclusion was derived, so others can check each reasoning "
         "step. Explored in research; no W3C standard."),
        ("Unifying logic", "One logical foundation under OWL, rules and queries, so they agree on "
         "what follows. A research goal, not a standard."),
        ("Crypto", "Signatures and encryption beside every layer, to show who published a graph and "
         "that it is unchanged. RDF Dataset Canonicalization (W3C, 2024) makes graphs signable."),
        ("XML", "An older interchange syntax. It lives on as RDF/XML, one way to write triples, and "
         "as XML Schema datatypes such as xsd:integer and xsd:date."),
    ]
    for i, (head, body) in enumerate(layers):
        y = 2.2 + i * 0.68
        P.circle(slide, x0 + 0.17, y + 0.17, 0.32, T.CLAY, "ABCDEF"[i], size=11)
        P.text(slide, x0 + 0.5, y, R - x0 - 0.5, 0.66,
               [[(head + "  ", {"bold": True}), (body, {"color": T.MUTED})]], size=12)
    P.takeaway(slide, x0, 6.35, R - x0,
               "Lower layers say what data means; upper layers would say whether to believe it.", size=14)
    P.notes(slide, """A quick tour of the layers this review skips. The user interface is whatever
        applications build on top. Trust, proof and unifying logic were the original vision for a Web
        where machines could check and weigh each other's conclusions; they remain mostly research,
        though W3C PROV standardizes provenance. Crypto runs beside every layer; RDF Dataset
        Canonicalization lets a graph be hashed and signed, as Verifiable Credentials do. XML was the
        original interchange syntax; today Turtle and JSON-LD are more common, but XML Schema
        datatypes are still how RDF types its literal values.""")
    return slide



# --------------------------------------------------------------------------- #
# Part 3: one slide per building block. Left: what it is and key terms.
# Right: a live example from the movie graph. Bottom: what it adds, what's missing.
# --------------------------------------------------------------------------- #
LX, LW = L, 5.95          # left column
RX = L + LW + 0.4         # right column
RW = R - RX


def s_layer_uri(prs, ex):
    slide = layer_frame(prs, 0, "Names: every thing gets a global identifier")
    what_it_is(slide, LX, 1.72, LW, 1.4,
               "A URI (Uniform Resource Identifier) is a Web-style name such as "
               "http://studio.example/JamesDean. An IRI is the same with any Unicode character allowed. "
               "In an ontology everything gets one: individuals, classes and properties alike.")
    terms(slide, LX, 3.2, LW, 2.9, [
        ("Global", "the same URI means the same thing in every file, store and query that uses it."),
        ("Namespace", "a URI prefix owned by one publisher, such as http://studio.example/."),
        ("Prefixed name", ":JamesDean is shorthand for the full URI once the prefix is declared."),
        ("Literal", "a plain value such as \"James Dean\" or 1955. It has no URI, but can carry a "
                    "datatype (xsd:integer) or a language tag (@en)."),
        ("Nonunique naming", "two different URIs can still mean the same thing; nothing forces "
                             "publishers to agree."),
    ], mono=False)
    prefixes = ("@prefix :    <http://studio.example/> .\n"
                "@prefix fan: <http://fanwiki.example/id/> .")
    P.code(slide, RX, 1.75, RW, 0.95, prefixes, label="Turtle", size=12)
    P.text(slide, RX, 2.85, RW, 0.3, "Every building block is named the same way", size=13, bold=True,
           color=T.MUTED)
    P.table(slide, RX, 3.15, RW, ("Block", "Written as", "Stands for"), [
        ("Individual", ":JamesDean", "http://studio.example/JamesDean"),
        ("Class", ":Actor", "http://studio.example/Actor"),
        ("Property", ":starring", "http://studio.example/starring"),
        ("Standard term", "rdfs:label", "http://www.w3.org/2000/01/rdf-schema#label"),
    ], col_widths=[1.3, 1.3, RW - 2.6], size=12, row_h=0.36, mono_cols=(1, 2))
    nodes = {"a": Node(RX + 1.0, 5.6, ":JamesDean", w=1.75), "b": Node(R - 1.15, 5.6, "fan:Dean_James", w=2.1)}
    P.graph(slide, nodes, [("a", "b", "same person?", "warn")], label_size=11)
    P.text(slide, RX, 5.95, RW, 0.25, "two publishers, two names: URIs alone can't tell (block 4 can)",
           size=11, italic=True, color=T.CLAY, align=PP_ALIGN.CENTER)
    adds_missing(slide, "a global name for every individual, class and property.",
                 "a name alone says nothing. RDF makes statements with names.")
    P.notes(slide, """Building block 1. Before you can say anything, things need names that work
        across systems. RDF borrows the Web's answer: URIs. A namespace groups names under one
        publisher, and prefixed names keep them short. Plain values are literals. Because publishers
        don't coordinate, one thing can have several URIs, which blocks 4 and 5 deal with.""")
    return slide


def s_layer_rdf(prs, ex):
    slide = layer_frame(prs, 1, "Statements: every fact is a triple")
    what_it_is(slide, LX, 1.72, LW, 1.5,
               "RDF (Resource Description Framework) writes all information as triples. The subject is "
               "the thing described, the predicate is the property, and the object is another thing or a "
               "plain value. Triples link up into a graph: things are nodes, properties are arrows.")
    terms(slide, LX, 3.3, LW, 2.85, [
        ("triple", "one statement: :Giant :starring :JamesDean ."),
        ("graph", "a set of triples; everything said about a thing meets at its node."),
        ("rdf:type", "says which class a thing belongs to; in Turtle it is the keyword a."),
        ("merge", "combining sources is the union of their triples; equal URIs join by themselves."),
        ("Turtle", "the syntax used here. A semicolon repeats the subject, a comma repeats subject "
                   "and predicate, a full stop ends. N-Triples, RDF/XML, JSON-LD: same triples."),
    ])
    P.code(slide, RX, 1.75, RW, 1.95, ex["turtle"].code, label="Turtle", size=12)
    P.text(slide, RX, 3.82, RW, 0.3, "The same block, read as triples", size=13, bold=True,
           color=T.MUTED)
    t = ex["table"]
    P.table(slide, RX, 4.12, RW, t.header, t.rows, col_widths=[2.05, 1.45, RW - 3.5], size=12,
            row_h=0.33, mono_cols=(0, 1, 2))
    adds_missing(slide, "individuals and the facts about them: the data.",
                 "to a machine, :Actor and :starring are just names. RDFS says what they mean.")
    P.notes(slide, """Building block 2. RDF breaks every fact into a triple. The Turtle on the
        right is the first movie in data/studio.ttl; the table lists the triples it contains, one per
        row. A set of triples is a graph, and merging graphs is just taking the union. RDF on its own
        is only data: nothing yet tells a machine what :starring or :Movie mean.""")
    return slide


def s_layer_rdfs(prs, ex):
    slide = layer_frame(prs, 2, "Vocabulary: classes, properties, hierarchies")
    what_it_is(slide, LX, 1.72, LW, 1.5,
               "RDF Schema is a small vocabulary, itself written as triples, for describing the classes "
               "and properties that data uses. Each term stands for an inference rule: when its pattern "
               "appears, new triples follow.")
    terms(slide, LX, 3.1, LW, 2.4, [
        ("rdfs:subClassOf", "every member of the subclass is in the superclass: every :Actor is a :Person."),
        ("rdfs:subPropertyOf", "what holds with the narrower property holds with the broader one."),
        ("rdfs:domain", "whatever has the property belongs to this class."),
        ("rdfs:range", "whatever the property points to belongs to this class."),
        ("rdfs:label", "a human-readable name; rdfs:comment adds a description."),
    ])
    box = P.rect(slide, LX, 5.45, LW, 0.6, fill=T.CLAY_TINT, radius=0.1, name="Caution")
    P.shape_text(box, [[("Careful: ", {"bold": True, "color": T.CLAY}),
                        ("domain and range add types; they never reject data. To validate data, use "
                         "SHACL (W3C, 2017).", {})]], size=12, align=PP_ALIGN.LEFT, margin=0.18)
    s = ex["subclass"]
    P.code(slide, RX, 1.75, RW, 2.3, s.code, label="Turtle", size=12)
    P.text(slide, RX, 4.17, RW, 0.3,
           f"Inferred by an RDFS reasoner: {len(s.rows)} of {s.facts['inferred_types']} new type triples",
           size=13, bold=True, color=T.MUTED)
    P.table(slide, RX, 4.47, RW, s.header, s.rows, col_widths=[2.0, 1.4, RW - 3.4], size=12, row_h=0.3,
            mono_cols=(0, 1, 2), header_fill=T.AMBER)
    adds_missing(slide, "classes, properties and hierarchies: the schema.",
                 "no way to say two names mean the same thing. RDFS-Plus adds that.")
    P.notes(slide, """Building block 3. RDFS gives meaning to class and property names through
        inference rules. :starring has range :Actor, so whoever is starred is an Actor; Actor is a
        subclass of Person, so they are a Person too. Domain and range are the classic beginner trap:
        they look like constraints but only classify. For validation, the later W3C standard SHACL
        checks data against shapes.""")
    return slide


def s_layer_rdfs_plus(prs, ex):
    slide = layer_frame(prs, 3, "Sameness: when two names mean one thing")
    what_it_is(slide, LX, 1.72, LW, 1.5,
               "RDFS-Plus is not a separate W3C standard. It is the small, practical slice of OWL that "
               "Semantic Web for the Working Ontologist teaches first: terms for aligning sources and for "
               "describing how properties behave.")
    terms(slide, LX, 3.1, LW, 3.0, [
        ("owl:equivalentClass", "two vocabularies' classes mean the same: fan:Star = :Actor."),
        ("owl:equivalentProperty", "the same for properties: fan:name = :personName."),
        ("owl:sameAs", "two URIs name one individual."),
        ("owl:inverseOf", ":starredIn is :starring read backwards."),
        ("owl:TransitiveProperty", "chains through: part of a part is a part."),
        ("owl:InverseFunctionalProperty", "a value identifies its subject, so sharing one implies sameAs."),
    ])
    sa = ex["same_as"]
    P.code(slide, RX, 1.75, RW, 1.75, sa.code, label="Turtle", size=11)
    nodes = {
        "fdean": Node(RX + 1.1, 3.95, "fan:Dean_James", w=2.1),
        "dean": Node(R - 0.95, 3.95, ":JamesDean", w=1.75),
        "pid": Node(RX + RW / 2, 5.0, '"P-0015"', kind="literal", w=1.2),
    }
    P.graph(slide, nodes, [
        ("fdean", "pid", "fan:registryId", "asserted", (-0.55, 0)),
        ("dean", "pid", ":registryId", "asserted", (0.5, 0)),
        ("fdean", "dean", "owl:sameAs", "inferred"),
    ], label_size=11)
    P.text(slide, RX, 5.45, RW, 0.65, [
        [("“Where was the star of Giant born?”", {"bold": True})],
        [(f"{sa.facts['before_rows']} rows before alignment  →  ", {"color": T.CLAY, "bold": True}),
         (f"“{sa.rows[0][1]}” after inference", {"color": T.AMBER, "bold": True})]], size=13,
        space_after=2)
    adds_missing(slide, "links between vocabularies and between identities.",
                 "can't define a class by conditions on its members. OWL can.")
    P.notes(slide, """Building block 4. Different publishers use different names. RDFS-Plus lets you
        align them: equivalentClass and equivalentProperty for vocabulary, sameAs for individuals.
        Declaring registryId inverse functional lets the reasoner work out sameness itself: two
        records with the same registry id are the same person, so the studio's movie data and the
        fan wiki's birthplace connect, and a question spanning both gets an answer.""")
    return slide


def s_layer_owl(prs, ex):
    slide = layer_frame(prs, 4, "Definitions: classes from conditions")
    what_it_is(slide, LX, 1.72, LW, 1.6,
               "OWL, the Web Ontology Language (OWL 2 since 2009), is the most expressive layer. Its "
               "axioms define a class by the values its members have, combine classes like sets, and "
               "count values. A reasoner then sorts individuals into classes and finds contradictions.")
    terms(slide, LX, 3.3, LW, 2.85, [
        ("owl:Restriction", "a class defined by a condition: hasValue, someValuesFrom, allValuesFrom."),
        ("owl:unionOf", "and intersectionOf, complementOf, disjointWith: classes as sets."),
        ("owl:maxCardinality", "and min, exact: count values. A movie has at most one director."),
        ("owl:propertyChainAxiom", "OWL 2: derive a property from a path of others; hasKey for ids."),
        ("Profiles", "DL, EL, QL, RL: subsets trading expressivity for guaranteed or fast reasoning."),
    ])
    hv = ex["has_value"]
    P.code(slide, RX, 1.75, RW, 1.5, hv.code, label="Turtle", size=12)
    P.text(slide, RX, 3.37, RW, 0.3, "Inferred by OWL RL: who belongs to :JamesDeanMovie", size=13,
           bold=True, color=T.MUTED)
    P.table(slide, RX, 3.67, RW, hv.header, hv.rows, col_widths=[2.4, 1.3, RW - 3.7], size=12,
            row_h=0.33, mono_cols=(0, 1, 2), header_fill=T.AMBER)
    box = P.rect(slide, RX, 5.2, RW, 0.85, fill=T.MIST, radius=0.1, name="Open world")
    P.shape_text(box, [[("Open world: ", {"bold": True}),
                        ("a missing fact is unknown, not false. Give a max-one-director movie two "
                         "directors and OWL concludes they are one person (Part 7).", {})]],
                 size=12, align=PP_ALIGN.LEFT, margin=0.18)
    adds_missing(slide, "formal definitions and constraints: the axioms.",
                 "some if-then logic fits no OWL pattern. Rules cover it.")
    P.notes(slide, """Building block 5. OWL lets you define classes instead of just naming them.
        JamesDeanMovie is defined as anything starring James Dean, and the reasoner finds the members
        on its own, from any source. OWL also has set operators and counting. Because the Web is an
        open world, OWL is cautious: it rarely concludes that something is false. Four profiles trade
        expressivity for efficient or guaranteed reasoning; owlrl implements OWL 2 RL.""")
    return slide


def s_layer_skos(prs, ex):
    slide = layer_frame(prs, 5, "Concepts: taxonomies and thesauri as linked data")
    what_it_is(slide, LX, 1.72, LW, 1.6,
               "SKOS (Simple Knowledge Organization System, W3C 2009) publishes controlled vocabularies, "
               "taxonomies and thesauri on the Web. Its unit is the concept: a subject you tag things "
               "with, rather than a class of things. It is built on RDFS-Plus.")
    terms(slide, LX, 3.3, LW, 2.1, [
        ("skos:Concept", "a subject term; skos:ConceptScheme is the vocabulary it belongs to."),
        ("skos:prefLabel", "the preferred name; skos:altLabel holds synonyms."),
        ("skos:broader", "and narrower, related: the hierarchy and cross-links."),
        ("skos:exactMatch", "and closeMatch, broadMatch: link to another vocabulary's concept."),
    ])
    box = P.rect(slide, LX, 5.4, LW, 0.65, fill=T.AMBER_TINT, radius=0.1, name="Concept or class")
    P.shape_text(box, [[("Concept or class? ", {"bold": True}),
                        ("A movie is tagged with g:Drama; it is not an instance of a Drama class. "
                         "Use SKOS for subjects and tags, classes for kinds of things.", {})]],
                 size=12, align=PP_ALIGN.LEFT, margin=0.18)
    P.code(slide, RX, 1.75, RW, 1.7, ex["skos"].facts["snippet"], label="Turtle", size=12)
    nodes = {
        "drama": Node(RX + 1.6, 4.0, "g:Drama", w=1.4),
        "lib": Node(R - 1.05, 4.0, "lib:DramaFilms", kind="ghost", w=1.95),
        "teen": Node(RX + 0.95, 5.3, "g:TeenDrama", w=1.65),
        "family": Node(RX + 3.25, 5.3, "g:FamilyDrama", w=1.85),
    }
    P.graph(slide, nodes, [
        ("drama", "lib", "exactMatch"),
        ("teen", "drama", "broader", "asserted", (-0.3, 0)),
        ("family", "drama", "broader", "asserted", (0.3, 0)),
    ], label_size=11)
    P.text(slide, R - 2.1, 4.35, 2.1, 0.5, "another library's vocabulary, linked not merged", size=11,
           italic=True, color=T.MUTED, align=PP_ALIGN.RIGHT)
    adds_missing(slide, "shared subject vocabularies for tagging, browsing and search.",
                 "deriving new relationships with if-then logic. Rules do that.")
    P.notes(slide, """Building block 6. Many organizations already have taxonomies and thesauri.
        SKOS publishes them as linked data: every concept gets a URI, labels carry synonyms, broader
        and narrower give the hierarchy, and matching properties link to other vocabularies without
        merging them. SKOS concepts are deliberately not classes: tagging a film with Drama doesn't
        make it an instance of a class, and broader is not transitive unless you choose.""")
    return slide


def s_layer_rules(prs, ex):
    slide = layer_frame(prs, 6, "Derivations: if-then rules beyond OWL")
    what_it_is(slide, LX, 1.72, LW, 1.9,
               "A rule says: IF this pattern is in the graph, THEN add that triple. RIF, the Rule "
               "Interchange Format (W3C, 2010), is the standard for exchanging rules between rule engines. "
               "In practice, rules are often written as SPARQL CONSTRUCT queries or in an engine's own "
               "rule language.")
    terms(slide, LX, 3.6, LW, 2.5, [
        ("Rules run the other layers", "RDFS and OWL RL are rule sets; owlrl, used for this deck, "
                                       "runs OWL as rules."),
        ("Use them to derive", "“worked with” = starred in a movie someone directed."),
        ("OWL 2 covers some rules", "property chains and keys express common rules as axioms."),
        ("Trade-off", "large rule sets are hard to maintain, and unrestricted rules carry no "
                      "guarantee of an answer."),
    ], heading="WHAT TO KNOW", mono=False)
    c, ch = ex["construct"], ex["chain"]
    P.text(slide, RX, 1.72, RW, 0.3, "AS A RULE: SPARQL CONSTRUCT", size=11, bold=True, color=T.TEAL)
    P.code(slide, RX, 2.02, RW, 1.65, c.code, label="SPARQL", size=11)
    P.text(slide, RX, 3.77, RW, 0.3, "AS AN OWL 2 AXIOM: PROPERTY CHAIN", size=11, bold=True, color=T.TEAL)
    P.code(slide, RX, 4.07, RW, 0.75, ch.code, label="Turtle", size=11)
    P.table(slide, RX, 4.95, RW, ch.header, ch.rows, col_widths=[1.75, 1.75, RW - 3.5], size=11,
            row_h=0.27, mono_cols=(0, 1, 2), header_fill=T.AMBER)
    adds_missing(slide, "derived relationships, written as if-then logic.",
                 "a way for applications to ask. SPARQL is the query language.")
    P.notes(slide, """Building block 7. Some knowledge is easiest to state as IF/THEN. Here "worked
        with" means starring in a movie that someone directed. The same three triples come out of a
        SPARQL CONSTRUCT rule and an OWL 2 property chain; the build checks that they agree. RIF is
        the W3C exchange format for rules, though many projects write rules in SPARQL or in their
        engine's language. Rules are flexible but large rule bases are hard to maintain.""")
    return slide


def s_layer_sparql(prs, ex):
    slide = layer_frame(prs, 7, "Questions: query asserted and inferred facts alike")
    what_it_is(slide, LX, 1.72, LW, 1.6,
               "SPARQL (W3C; version 1.1 since 2013) is the query language for RDF. A query is a graph "
               "pattern: triples with ?variables standing in for the unknowns. The engine returns every "
               "way the pattern matches the graph.")
    terms(slide, LX, 3.3, LW, 2.85, [
        ("SELECT", "returns a table of variable bindings, as in SQL."),
        ("ASK", "returns yes or no; handy for testing a model."),
        ("CONSTRUCT", "builds new triples from a template, so it can act as a rule."),
        ("DESCRIBE", "returns what the store knows about a resource."),
        ("skos:broader*", "a property path: follow a link any number of times."),
        ("SERVICE", "sends part of a query to a remote SPARQL endpoint."),
    ])
    q = ex["select"]
    P.code(slide, RX, 1.75, RW, 2.6, q.code, label="SPARQL", size=12)
    P.text(slide, RX, 4.47, RW, 0.3, "Result over data/studio.ttl", size=13, bold=True, color=T.MUTED)
    P.table(slide, RX, 4.77, RW, q.header, q.rows, col_widths=[2.6, 0.8, RW - 3.4], size=12, row_h=0.31)
    adds_missing(slide, "a standard way for applications to ask the model and the data.",
                 "all eight blocks answering one question together.", missing_label="Next: ")
    P.notes(slide, """Building block 8. SPARQL is how applications use an ontology. A query looks
        like the data it targets, with variables for the unknowns; shared variables join triples
        without a JOIN clause. Run over a graph that a reasoner has expanded, a query sees what the
        ontology implies as well as what was typed in. The next slide shows that with one question
        that needs every layer.""")
    return slide


def s_together_flow(prs, ex):
    slide = content(prs, "Part 3 · The building blocks together",
                    "Merge, reason, query: the layers as one system")
    f = ex["together"].facts
    steps = [
        ("Sources", f"Studio catalogue and fan wiki: {f['data']} triples, two names for James Dean"),
        ("Merge", "URI + RDF: parse into triples and take the union; equal URIs join"),
        ("Reason", f"Rules from RDFS, RDFS-Plus and OWL add {f['inferred']} triples nobody typed in"),
        ("Query", "SPARQL matches one pattern over asserted and inferred triples"),
        ("Answer", "James Dean, born in Marion, Indiana, worked with three directors"),
    ]
    w, gap = 2.1, (CW - 5 * 2.1) / 4
    for i, (head, body) in enumerate(steps):
        x = L + i * (w + gap)
        hot = head == "Reason"
        P.rect(slide, x, 1.85, w, 2.0, fill=T.TEAL if hot else T.MIST, radius=0.12, name=f"Step: {head}")
        P.text(slide, x + 0.18, 2.0, w - 0.36, 0.6, head, size=16, bold=True,
               color=T.WHITE if hot else T.INK)
        P.text(slide, x + 0.18, 2.5, w - 0.36, 1.3, body, size=13, color=T.WHITE if hot else T.INK)
        if i < 4:
            P.line(slide, x + w + 0.05, 2.85, x + w + gap - 0.05, 2.85, color=T.TEAL, width=2)
    sx = L + 2 * (w + gap) + w / 2
    P.line(slide, sx, 3.9, sx, 4.3, color=T.AMBER, width=2, dashed=True, arrow=False)
    box = P.rect(slide, sx - 2.6, 4.3, 5.2, 0.85, fill=T.AMBER_TINT, line=T.AMBER, radius=0.1)
    P.shape_text(box, f"The ontology: {f['model']} triples of RDFS, RDFS-Plus, OWL and SKOS, "
                      "stored in the same graph as the data", size=14)
    P.takeaway(slide, L, 5.45, CW,
               "The data says what happened; the ontology says what it means; reasoning and SPARQL "
               "turn both into answers.", size=14)
    pw = (CW - 7 * 0.15) / 8
    for i, (name, role) in enumerate(LAYERS):
        box = P.rect(slide, L + i * (pw + 0.15), 6.1, pw, 0.6, fill=T.WHITE, line=T.TEAL, radius=0.3,
                     name=f"Layer: {name}")
        P.shape_text(box, [[(name, {"bold": True, "size": 12})], [(role, {"size": 10, "color": T.MUTED})]],
                     color=T.TEAL)
    P.notes(slide, f"""Putting the blocks back together, in the same shape as the application
        architecture in Part 4. Sources are converted to triples and merged ({f['data']} triples of
        data). The ontology, {f['model']} more triples, sits in the same store. A reasoner applies
        its rules and adds {f['inferred']} triples about the example. SPARQL then asks one question
        across all of it. The next slide shows the question line by line.""")
    return slide


def s_together_query(prs, ex):
    slide = content(prs, "Part 3 · The building blocks together",
                    "One question that needs every layer")
    t = ex["together"]
    P.text(slide, L, 1.7, CW, 0.4,
           "“Where were the stars of drama films born, and which directors did they work with?”",
           size=17, bold=True, color=T.TEAL_DARK)
    why = {
        "SKOS": "Teen and family drama are narrower than g:Drama",
        "RDF": "Asserted by the studio catalogue",
        "RDFS": "From range :Actor, and every Actor is a Person",
        "RDFS-Plus": "fan:Dean_James sameAs :JamesDean, via registry id",
        "OWL": "Property chain: starredIn, then directedBy",
    }
    label_why = "From :personName and fan:name, sub-properties of rdfs:label"
    rows, fills = [], {}
    for r, (pattern, layer) in enumerate(t.facts["lines"]):
        reason = label_why if "rdfs:label" in pattern else why[layer]
        rows.append((pattern, layer, reason))
        fills[(r, 1)] = T.MIST if layer == "RDF" else T.AMBER_TINT
    P.text(slide, L, 2.2, 7.4, 0.3, "THE QUERY'S PATTERN, LINE BY LINE", size=11, bold=True, color=T.TEAL)
    P.table(slide, L, 2.5, 7.4, ("Pattern line", "Layer", "Why it matches"), rows,
            col_widths=[3.6, 1.0, 2.8], size=12, row_h=0.47, mono_cols=(0,), cell_fills=fills)
    x0 = L + 7.75
    P.text(slide, x0, 2.2, R - x0, 0.3, "RESULT (OWL RL REASONER + SPARQL)", size=11, bold=True,
           color=T.TEAL)
    P.table(slide, x0, 2.5, R - x0, t.header, t.rows, col_widths=[1.1, 1.65, R - x0 - 2.75], size=12,
            row_h=0.34, header_fill=T.AMBER)
    P.text(slide, x0, 4.05, R - x0, 0.3, "TAKE ONE LAYER AWAY", size=11, bold=True, color=T.TEAL)
    abl = [(label, str(n)) for label, n in t.facts["ablations"]]
    P.table(slide, x0, 4.35, R - x0, ("Run", "Rows"), abl, col_widths=[R - x0 - 0.8, 0.8], size=11,
            row_h=0.3, cell_colors={(r, 1): T.CLAY for r in range(len(abl))})
    P.takeaway(slide, x0, 6.3, R - x0,
               "Every layer is load-bearing: drop any one and the answer disappears.", size=13)
    P.notes(slide, """The only line that matches asserted data is the starring triple; every other
        line depends on a layer of the ontology. The table on the right is the live result, and the
        ablation reruns the same query with one file removed at a time, or with no reasoner. Each
        run returns zero rows, which the build checks. Query: queries/together.rq.""")
    return slide


def s_cheatsheet(prs, ex):
    slide = content(prs, "Part 3 · The building blocks together", "Cheat sheet: the building blocks at a glance")
    rows = [
        ("1 · URI", "What is this thing called, globally?", "URI, IRI, namespace, prefix, literal",
         ":JamesDean"),
        ("2 · RDF", "What is true about it?", "triple, graph, rdf:type, Turtle, merge",
         ":Giant :starring :JamesDean"),
        ("3 · RDFS", "What kinds of things and relations exist?",
         "subClassOf, subPropertyOf, domain, range, label", ":Actor rdfs:subClassOf :Person"),
        ("4 · RDFS-Plus", "Which names mean the same thing?",
         "sameAs, equivalentClass, inverseOf, InverseFunctionalProperty", "same registry id ⇒ owl:sameAs"),
        ("5 · OWL", "What exactly defines a class?", "Restriction, hasValue, unionOf, cardinality, profiles",
         ":JamesDeanMovie = starring :JamesDean"),
        ("6 · SKOS", "How are subject terms organized?", "Concept, prefLabel, broader, exactMatch",
         "g:TeenDrama skos:broader g:Drama"),
        ("7 · Rules", "What else follows, if-then?", "rule, CONSTRUCT, property chain, RIF",
         "starredIn + directedBy ⇒ workedWith"),
        ("8 · SPARQL", "How do I ask?", "SELECT, ASK, CONSTRUCT, graph pattern, property path",
         "?movie :starring :JamesDean"),
    ]
    P.table(slide, L, 1.75, CW, ("Block", "Answers the question", "Key terms", "In the movie example"),
            rows, col_widths=[1.6, 3.2, 3.85, CW - 8.65], size=12, row_h=0.52, mono_cols=(3,),
            cell_fills={(r, 0): T.MIST for r in range(len(rows))})
    P.takeaway(slide, L, 6.55, CW,
               "Model only as far as your questions need: RDFS-Plus covers most integration work.", size=14)
    P.notes(slide, """A one-page reference for the eight building blocks: what question each layer
        answers, the terms to recognize, and where it shows up in the movie example. The deep-dive
        parts that follow go through each layer in more detail.""")
    return slide


# --------------------------------------------------------------------------- #
# Part 2: why we need ontologies (continued)
# --------------------------------------------------------------------------- #
def s_aaa(prs, ex):
    slide = content(prs, "Part 2 · Why the Web makes it harder",
                    "Anyone can say Anything about Any topic")
    P.text(slide, L, 1.7, CW, 0.5,
           "The Semantic Web links data, not pages: facts about one entity can live on many "
           "servers, joined by global URIs.", size=16, color=T.MUTED)
    w = (CW - 0.6) / 3
    cards = [
        ("Open world", "There is always more someone could say. A missing fact is unknown, not false."),
        ("Nonunique naming", "Publishers don't coordinate, so one entity goes by many names."),
        ("Network effect", "Each new publisher adds value for all, if the infrastructure copes "
                           "with messy, varied data."),
    ]
    for i, (head, body) in enumerate(cards):
        P.card(slide, L + i * (w + 0.3), 2.4, w, 2.0, head, body, badge=i + 1)
    P.rect(slide, L, 4.7, CW, 1.95, fill=None, line=T.LINE, radius=0.14)
    P.text(slide, L + 0.3, 4.9, 4.0, 0.4, "Running example: two sources, one actor", size=16, bold=True)
    P.text(slide, L + 0.3, 5.35, 4.0, 1.1,
           "A studio catalogue and a fan wiki both describe James Dean, under different names.",
           size=14, color=T.MUTED)
    nodes = {
        "a": Node(6.3, 5.67, ":JamesDean", w=2.0),
        "b": Node(10.9, 5.67, "fan:Dean_James", w=2.3),
    }
    P.graph(slide, nodes, [("a", "b", "same person?", "warn")])
    P.text(slide, 5.3, 6.08, 2.0, 0.3, "studio.example", size=11, color=T.MUTED, align=PP_ALIGN.CENTER)
    P.text(slide, 9.75, 6.08, 2.3, 0.3, "fanwiki.example", size=11, color=T.MUTED, align=PP_ALIGN.CENTER)
    P.notes(slide, """The AAA slogan: because publication is in the hands of whoever produces
        the content, Web data is a wilderness, valuable but tangled. Two consequences: assume an
        open world, and expect names to be nonunique. An ontologist's job is to turn this mess into
        durable information resources using RDF, RDFS, SPARQL, SKOS and OWL.""")
    return slide


def s_modeling(prs, ex):
    slide = content(prs, "Part 2 · Semantic modeling",
                    "Modeling: organizing information for community use")
    roles = [
        ("Communicate", "Share one picture of a situation."),
        ("Explain and predict", "Relate phenomena to one another."),
        ("Mediate", "Capture what people agree on; leave room where they differ."),
    ]
    for i, (head, body) in enumerate(roles):
        P.card(slide, L, 1.8 + i * 1.6, 5.4, 1.4, head, body, badge=i + 1)
    x0 = 6.6
    P.text(slide, x0, 1.8, R - x0, 0.4, "Structured and unstructured knowledge always coexist",
           size=17, bold=True)
    shades = [T.MIST, T.MINT_SOFT, T.MINT, T.TEAL, T.INK]
    sw = (R - x0 - 0.4) / 5
    for i, fill in enumerate(shades):
        P.rect(slide, x0 + i * (sw + 0.1), 2.45, sw, 0.75, fill=fill, radius=0.08)
    P.text(slide, x0, 3.3, 2.8, 0.6, "Discordant: everyone describes it their own way", size=13,
           color=T.MUTED)
    P.text(slide, R - 2.8, 3.3, 2.8, 0.6, "Formal: agreed and machine-readable", size=13,
           color=T.MUTED, align=PP_ALIGN.RIGHT)
    P.line(slide, x0, 4.1, R, 4.1, color=T.AMBER, width=2.5)
    P.text(slide, x0, 4.18, R - x0, 0.35, "continuous distillation, never “done”", size=13,
           italic=True, color=T.AMBER, align=PP_ALIGN.CENTER)
    half = (R - x0 - 0.2) / 2
    P.pill(slide, x0, 4.75, half, 0.45, "Commonality: what things share", T.TEAL, size=13)
    P.pill(slide, x0 + half + 0.2, 4.75, half, 0.45, "Variability: how they differ", T.TEAL, size=13)
    box = P.rect(slide, x0, 5.5, R - x0, 0.95, fill=T.AMBER_TINT, radius=0.12)
    P.shape_text(box, "Standards are a medium for collaborating on models, not just for sharing data.",
                 size=15, italic=True, align=PP_ALIGN.LEFT, margin=0.25)
    P.notes(slide, """Models let people communicate, explain and predict, and mediate between
        viewpoints. Structured and unstructured knowledge about a topic coexist forever, so the
        job is a continuous process of distilling communal knowledge without forcing agreement
        up front. Any modeling language must capture commonality while allowing variability.""")
    return slide


# --------------------------------------------------------------------------- #
# Part 4: deep dive, RDF and SPARQL
# --------------------------------------------------------------------------- #
def s_table_to_triples(prs, ex):
    slide = content(prs, "Part 4 · RDF", "From a table to triples")
    P.text(slide, L, 1.7, CW, 0.45,
           "Spread a table across servers by row, by column, or by cell? Only cell by cell lets "
           "anyone say anything about any entity, so every cell becomes a triple.",
           size=16, color=T.MUTED)
    for i, (label, ok) in enumerate([("by row", False), ("by column", False), ("by cell", True)]):
        P.pill(slide, L + i * 1.75, 2.5, 1.6, 0.4, label + ("  ✓" if ok else "  ✗"),
               T.TEAL if ok else T.WHITE, color=T.WHITE if ok else T.MUTED, size=13,
               line=None if ok else T.LINE)
    header = ("", "movieTitle", "releaseYear", "directedBy")
    rows = [(":Rebel…", "Rebel Without a Cause", "1955", ":NicholasRay"),
            (":EastOfEden", "East of Eden", "1955", ":EliaKazan")]
    P.table(slide, L, 3.2, 5.9, header, rows, col_widths=[1.3, 2.1, 1.15, 1.35], size=13,
            row_h=0.42, cell_fills={(0, 3): T.AMBER_TINT})
    P.text(slide, L, 4.65, 5.9, 0.9, [
        [("row", {"bold": True, "color": T.TEAL}), (" → subject   ", {}),
         ("column", {"bold": True, "color": T.TEAL}), (" → predicate   ", {}),
         ("cell", {"bold": True, "color": T.AMBER}), (" → object", {})],
    ], size=15)
    P.line(slide, 6.6, 3.85, 7.05, 3.85, color=T.AMBER, width=2.5)
    triples = ex["table"]
    P.table(slide, 7.15, 3.2, R - 7.15, triples.header, triples.rows, col_widths=[2.1, 1.45, 2.03],
            size=12, row_h=0.38, mono_cols=(0, 1, 2),
            cell_fills={(2, 1): T.AMBER_TINT, (2, 2): T.AMBER_TINT})
    P.takeaway(slide, L, 6.15, CW,
               "No columns to line up and no nulls: a relationship is either stated or it isn't.")
    P.notes(slide, """Start from a plain table. Only cell-by-cell distribution honours
        AAA, so each cell becomes a triple: subject is the row, predicate the column, object the
        value. The highlighted cell, directed by Nicholas Ray, is the highlighted triple. The
        triples on the right are read from data/studio.ttl at build time.""")
    return slide


def s_merge(prs, ex):
    slide = content(prs, "Part 4 · RDF", "Merging graphs is just a union of triples")
    m = ex["merge"].facts
    P.text(slide, L, 1.75, 4, 0.3, "studio.example", size=12, bold=True, color=T.TEAL)
    P.text(slide, 6.55, 1.75, 2.1, 0.3, "fanwiki.example", size=12, bold=True, color=T.TEAL,
           align=PP_ALIGN.CENTER)
    nodes = {
        "rebel": Node(2.05, 2.55, ":RebelWithoutACause", w=2.75),
        "giant": Node(2.05, 4.85, ":Giant", w=1.3),
        "dean": Node(4.4, 3.7, ":JamesDean", w=1.85),
        "fdean": Node(7.6, 2.55, "fan:Dean_James", w=2.1),
        "pid": Node(5.9, 5.75, '"P-0015"', kind="literal", w=1.3),
        "born": Node(8.25, 4.45, '"Marion, Indiana"', kind="literal", w=1.9),
    }
    P.graph(slide, nodes, [
        ("rebel", "dean", ":starring", "asserted", (0.25, -0.1)),
        ("giant", "dean", ":starring", "asserted", (0.25, 0.1)),
        ("dean", "pid", ":registryId", "asserted", (-0.5, 0.0)),
        ("fdean", "pid", "fan:registryId", "asserted", (-0.35, 0.0)),
        ("fdean", "born", "fan:bornIn", "asserted", (0.35, 0.0)),
        ("dean", "fdean", "same?", "warn", (0, -0.05)),
    ])
    x0 = 9.6
    P.stat(slide, x0, 1.8, R - x0, f"{m['studio']} + {m['fan']} = {m['merged']}",
           "triples: studio + fan wiki = merged graph", value_pt=36)
    P.text(slide, x0, 3.4, R - x0, 1.0,
           "No schema to align first: the merged store is the set union of both sources.",
           size=15)
    P.card(slide, x0, 4.5, R - x0, 2.1, "Identity is what's left",
           f"{m['dean_nodes'][0]} and {m['dean_nodes'][1]} share a literal but stay two nodes. "
           "Only a shared URI makes one node.",
           fill=T.CLAY_TINT, title_size=16, body_size=14)
    P.notes(slide, f"""Seen as a graph, each triple is a labelled edge and everything known about
        an entity gathers at one node. Merging sources reduces to taking the union: {m['studio']}
        studio triples plus {m['fan']} fan-wiki triples gives {m['merged']}. The flexibility
        moves the difficulty into identity: without a shared URI, the two James Dean nodes stay
        apart. RDFS-Plus fixes that in Part 5 with an inverse-functional property.""")
    return slide


def s_identity(prs, ex):
    slide = content(prs, "Part 4 · RDF", "Global identity with URIs, written in Turtle")
    P.code(slide, L, 1.8, 6.4, 2.15, ex["turtle"].code, label="Turtle")
    P.text(slide, L, 4.35, 6.4, 0.35, "A qname is shorthand for a full URI", size=15, bold=True)
    P.pill(slide, L, 4.85, 1.9, 0.5, ":JamesDean", T.TEAL, size=14, font=T.CODE_FONT)
    P.line(slide, L + 2.0, 5.1, L + 2.55, 5.1, color=T.TEAL, width=2)
    P.pill(slide, L + 2.65, 4.85, 3.75, 0.5, "<http://studio.example/JamesDean>", T.WHITE,
           color=T.TEAL, size=13, line=T.TEAL, font=T.CODE_FONT)
    P.text(slide, L, 5.55, 6.4, 0.5, "Same URI = same node, in every graph that mentions it.",
           size=14, color=T.MUTED)
    P.text(slide, 7.4, 1.8, R - 7.4, 4.9, [
        [("URIs are global names. ", {"bold": True}), ("Two nodes are the same exactly when they share one.", {})],
        [("Namespaces ", {"bold": True}), ("group names under one authority; one triple can mix several.", {})],
        [("rdf:type ", {"bold": True, "font": T.CODE_FONT}), ("links an instance to its type; ", {}),
         ("rdf:Property", {"bold": True, "font": T.CODE_FONT}), (" marks a predicate.", {})],
        [("Blank nodes ", {"bold": True}), ("stand for things that exist but have no global name.", {})],
        [("Turtle, N-Triples, RDF/XML ", {"bold": True}), ("are serializations of the same triples.", {})],
    ], size=16, bullets=True, space_after=14)
    P.notes(slide, """RDF borrows the Web's answer to identity: URIs, grouped in namespaces and
        abbreviated as qnames. Built-in rdf terms give a minimal type system. Blank nodes cover
        anonymous entities. The Turtle on the left is the first block of data/studio.ttl.""")
    return slide


def s_architecture(prs, ex):
    slide = content(prs, "Part 4 · Application architecture",
                    "Parse, store, query, then build the application")
    steps = [
        ("Sources", "Turtle, RDF/XML, spreadsheets, RDFa pages"),
        ("Parser and serializer", "Text ⇄ triples, plus converters for tables"),
        ("RDF store", "Merges sources: same URI = same node"),
        ("Query engine", "SPARQL, also served as web endpoints"),
        ("Application", "Analysis, UI, archiving; reaches data only via queries"),
    ]
    w, gap = 2.1, (CW - 5 * 2.1) / 4
    for i, (head, body) in enumerate(steps):
        x = L + i * (w + gap)
        hot = head == "RDF store"
        P.rect(slide, x, 1.85, w, 2.0, fill=T.TEAL if hot else T.MIST, radius=0.12, name=f"Step: {head}")
        P.text(slide, x + 0.18, 2.0, w - 0.36, 0.6, head, size=16, bold=True,
               color=T.WHITE if hot else T.INK)
        P.text(slide, x + 0.18, 2.65, w - 0.36, 1.15, body, size=14, color=T.WHITE if hot else T.INK)
        if i < 4:
            P.line(slide, x + w + 0.05, 2.85, x + w + gap - 0.05, 2.85, color=T.TEAL, width=2)
    sx = L + 2 * (w + gap) + w / 2
    P.line(slide, sx, 3.9, sx, 4.3, color=T.AMBER, width=2, dashed=True, arrow=False)
    box = P.rect(slide, sx - 2.0, 4.3, 4.0, 0.75, fill=T.AMBER_TINT, line=T.AMBER, radius=0.1,
                 dash=None)
    P.shape_text(box, "Models (RDFS, OWL) sit in the same store: they are triples too", size=14)
    P.takeaway(slide, L, 5.35, CW,
               "Federation: convert every source to triples, merge, then query, without "
               "caring where a triple came from.")
    uses = ["Calendar integration", "Map integration", "Annotation", "Content management"]
    uw = (CW - 3 * 0.3) / 4
    for i, label in enumerate(uses):
        P.pill(slide, L + i * (uw + 0.3), 6.05, uw, 0.45, label, T.WHITE, color=T.TEAL, size=14,
               line=T.TEAL)
    P.notes(slide, """A Semantic Web application chains a parser and serializer, an RDF store
        whose distinguishing feature is merging, a query engine, and application code. These
        handle the mechanics but not the meaning: semantic models supply meaning, and they live
        in the same store. Typical federation apps: calendars, maps, annotation, content
        management.""")
    return slide


def s_sparql_select(prs, ex):
    slide = content(prs, "Part 4 · SPARQL", "SPARQL: the question looks like the data")
    q = ex["select"]
    P.code(slide, L, 1.8, 6.3, 3.0, q.code, label="SPARQL")
    P.text(slide, L, 5.1, 6.3, 0.3, "The graph pattern inside WHERE", size=13, bold=True,
           color=T.MUTED)
    nodes = {
        "dir": Node(1.45, 6.05, "?director", kind="var", w=1.6, h=0.45),
        "m": Node(3.85, 6.05, "?movie", kind="var", w=1.3, h=0.45),
        "jd": Node(6.0, 6.05, ":JamesDean", w=1.6, h=0.45),
    }
    P.graph(slide, nodes, [("m", "dir", ":directedBy", "asserted", (0, -0.3)),
                           ("m", "jd", ":starring", "asserted", (0, -0.3))], label_size=11)
    x0 = 7.3
    P.text(slide, x0, 1.8, R - x0, 0.35, "Result over data/studio.ttl", size=15, bold=True)
    P.table(slide, x0, 2.25, R - x0, q.header, q.rows, col_widths=[2.45, 0.85, 2.13], size=14,
            row_h=0.42)
    P.text(slide, x0, 4.2, R - x0, 2.5, [
        "Variables stand in for the unknowns",
        "Shared variables join triples: no JOIN clause",
        "Predicates can be variables too",
        "Easy for SQL users: syntax borrows from SQL and Turtle",
    ], size=15, bullets=True, space_after=10)
    P.notes(slide, """RDF is a tell-and-ask system. SPARQL's answer to querying data with an
        unknown structure is the graph pattern: a small graph mixing known resources and
        variables, matched against the data. The query is queries/dean_movies.rq and the
        table is its live result.""")
    return slide


def s_sparql_explore(prs, ex):
    slide = content(prs, "Part 4 · SPARQL", "Explore unknown data, then build new graphs")
    half = (CW - 0.4) / 2
    xr = L + half + 0.4
    e, c = ex["explore"], ex["construct"]
    P.text(slide, L, 1.75, half, 0.4, "Explore: what is said here?", size=17, bold=True)
    P.code(slide, L, 2.25, half, 1.5, e.code, label="SPARQL")
    P.text(slide, L, 3.95, half, 0.3, "Result on the fan wiki", size=13, bold=True, color=T.MUTED)
    px = L
    for (p,) in e.rows:
        pw = 0.1 * len(p) + 0.4
        P.pill(slide, px, 4.35, pw, 0.42, p, T.WHITE, color=T.TEAL, size=12, line=T.TEAL,
               font=T.CODE_FONT)
        px += pw + 0.12
    P.text(slide, L, 5.05, half, 1.2,
           "A variable in the predicate position lists the properties of a data set you have "
           "never seen, so you can discover its schema.", size=14, color=T.MUTED)
    P.text(slide, xr, 1.75, half, 0.4, "CONSTRUCT: a query that acts as a rule", size=17, bold=True)
    P.code(slide, xr, 2.25, half, 2.15, c.code, label="SPARQL")
    P.table(slide, xr, 4.6, half, c.header, c.rows, col_widths=[1.75, 1.7, half - 3.45], size=12,
            row_h=0.36, mono_cols=(0, 1, 2))
    P.text(slide, xr, 6.2, half, 0.5, "Part 8 gets the same triples from an OWL 2 property chain.",
           size=13, color=T.MUTED)
    P.notes(slide, f"""Because variables can match predicates, one query reveals what an
        unfamiliar data set talks about. CONSTRUCT uses the same pattern matching to build a new
        graph from a template, which lets a query transform data and act as a rule. Here it
        builds {c.facts['total']} workedWith triples; the three for James Dean are shown.""")
    return slide


# --------------------------------------------------------------------------- #
# Part 5: deep dive, inference
# --------------------------------------------------------------------------- #
def s_inference(prs, ex):
    slide = content(prs, "Part 5 · RDF and inferencing", "Inference: stated facts imply new ones")
    s = ex["subclass"].facts
    box = P.rect(slide, 3.25, 1.85, 3.7, 0.45, fill=T.MIST, radius=0.08)
    P.shape_text(box, "schema: :starring rdfs:range :Actor", size=11, font=T.CODE_FONT, color=T.INK)
    nodes = {
        "giant": Node(1.85, 2.6, ":Giant", w=1.4),
        "dean": Node(1.85, 4.45, ":JamesDean", w=1.9),
        "actor": Node(5.6, 3.55, ":Actor", kind="class", w=1.4),
        "person": Node(5.6, 5.75, ":Person", kind="class", w=1.5),
    }
    P.graph(slide, nodes, [
        ("giant", "dean", ":starring"),
        ("dean", "actor", "rdf:type", "inferred", (0, -0.05)),
        ("actor", "person", "rdfs:subClassOf", "asserted", (0.05, 0)),
        ("dean", "person", "rdf:type", "inferred", (-0.15, 0.15)),
    ])
    P.text(slide, L, 6.15, 3.0, 0.4, "two inference steps, chained", size=13, italic=True, color=T.AMBER)
    x0 = 7.4
    rule = ("# type propagation (rdfs9)\nCONSTRUCT { ?x a ?B }\nWHERE {\n"
            "  ?A rdfs:subClassOf ?B .\n  ?x a ?A .\n}")
    P.code(slide, x0, 1.85, R - x0, 1.85, rule, label="Rule")
    P.stat(slide, x0, 4.0, 2.5, str(s["asserted"]), "asserted triples (data + schema)")
    P.stat(slide, x0 + 2.7, 4.0, 2.6, f"+{s['inferred_types']}", "inferred rdf:type triples",
           color=T.AMBER)
    P.text(slide, x0, 5.6, R - x0, 0.8,
           "Inferred triples differ from asserted ones only in provenance: the engine treats them the same.",
           size=14)
    P.notes(slide, f"""In the Semantic Web, a modeling term means exactly the inferences it
        licenses. subClassOf's whole meaning is type propagation, an IF/THEN rule that can
        be written as a SPARQL CONSTRUCT. Running RDFS inference over the studio data and schema
        ({s['asserted']} asserted triples) adds {s['inferred_types']} type triples, for example
        James Dean becomes an Actor via the range of starring, then a Person via subClassOf.""")
    return slide


def s_when(prs, ex):
    slide = content(prs, "Part 5 · RDF and inferencing", "When are inferences computed?")
    P.text(slide, L, 1.7, CW, 0.45, "Standards define which inferences are valid, not when they are computed.",
           size=16, color=T.MUTED)
    w = (CW - 0.6) / 3
    plus = {"bold": True, "color": T.TEAL}
    minus = {"bold": True, "color": T.CLAY}
    cols = [
        ("Cached", "Inferred triples are stored with the data.",
         [("+ ", plus, "Simple, fast queries"),
          ("– ", minus, "Triple explosion"),
          ("– ", minus, "Change management: which inferred triples go when a source triple is removed?")]),
        ("Just in time", "Computed at query time, never stored.",
         [("+ ", plus, "Always consistent with the sources"),
          ("– ", minus, "Repeats the same work on every query")]),
        ("No explicit inferencing", "Queries spell out what they need.",
         [("+ ", plus, "No inference engine to run"),
          ("– ", minus, "Every query must know the model; its meaning still lives in the rules")]),
    ]
    for i, (head, lead, items) in enumerate(cols):
        x = L + i * (w + 0.3)
        P.rect(slide, x, 2.4, w, 2.8, fill=T.MIST, radius=0.14, name=f"Card: {head}")
        P.text(slide, x + 0.25, 2.6, w - 0.5, 0.4, head, size=T.CARD_TITLE_PT, bold=True)
        P.text(slide, x + 0.25, 3.05, w - 0.5, 0.5, lead, size=14, color=T.MUTED)
        P.text(slide, x + 0.25, 3.65, w - 0.5, 1.8,
               [[(sign, style), (body, {})] for sign, style, body in items], size=14, space_after=8)
    P.text(slide, L, 5.55, CW, 0.35, "The languages differ only in which rules they include",
           size=15, bold=True)
    ladder = ["RDFS: classes and properties", "RDFS-Plus: + sameness and property traits",
              "OWL: + classes defined by property values"]
    for i, label in enumerate(ladder):
        P.pill(slide, L + i * (w + 0.3), 6.0, w, 0.45, label, [T.MINT, T.TEAL, T.INK][i],
               color=T.INK if i == 0 else T.WHITE, size=13)
    P.notes(slide, """Cached inferencing stores inferred triples with the data: simple, but it
        can explode and makes change management hard, since a triple may be inferable in
        several ways. Just-in-time inferencing computes at query time. A third option does no
        explicit inferencing, yet the inference reading still explains the model. RDFS,
        RDFS-Plus and OWL differ only in their rule sets.""")
    return slide


def s_rdfs(prs, ex):
    slide = content(prs, "Part 5 · RDF Schema", "RDFS: a schema made of inference rules")
    P.code(slide, L, 1.8, 6.4, 2.75, ex["subclass"].code, label="Turtle", size=12)
    P.text(slide, L, 4.85, 6.4, 1.9, [
        "Keywords like rdfs:subClassOf are ordinary resources; inference rules give them meaning",
        "Common superclass ≈ union; one subPropertyOf triple maps one property onto another",
        "Tempted to find-and-replace merged data? Use these constructs instead",
    ], size=14, bullets=True, space_after=6)
    x0 = 7.4
    lab = ex["labels"]
    P.text(slide, x0, 1.8, R - x0, 0.35, "?thing rdfs:label ?label, after RDFS inference", size=14,
           bold=True)
    movies = {r for r, row in enumerate(lab.rows) if row[0] in
              (":EastOfEden", ":Giant", ":RebelWithoutACause")}
    P.table(slide, x0, 2.25, R - x0, lab.header, lab.rows, col_widths=[2.6, R - x0 - 2.6], size=13,
            row_h=0.33, mono_cols=(0,), cell_fills={(r, 1): T.AMBER_TINT for r in movies})
    P.text(slide, x0, 6.0, R - x0, 0.7, [[("Movies", {"bold": True}),
           (" (from :movieTitle) and people (from :personName) answer one query.", {})]], size=13)
    P.notes(slide, """RDFS is written in RDF, so schema is as open to AAA as data. Its terms are
        defined by inference patterns, not validation. With movieTitle and personName both
        declared subproperties of rdfs:label, one query labels every movie and person: ten rows,
        none of which were asserted with rdfs:label.""")
    return slide


def s_domain_trap(prs, ex):
    slide = content(prs, "Part 5 · RDF Schema", "Domain and range classify, they never reject")
    P.text(slide, L, 1.75, 6.4, 0.35, "Source C: a theatre archive reuses :directedBy", size=15, bold=True)
    P.code(slide, L, 2.2, 6.4, 1.25, ex["domain_trap"].code, label="Turtle")
    P.text(slide, L, 3.65, 6.4, 0.35, "…merged with the studio schema", size=15, bold=True)
    P.code(slide, L, 4.1, 6.4, 0.7, ":directedBy rdfs:domain :Movie .", label="Turtle")
    nodes = {
        "play": Node(10.05, 2.35, "th:StreetcarNamedDesire1947", kind="warn", w=3.7),
        "movie": Node(10.05, 4.35, ":Movie", kind="class", w=1.4),
    }
    P.graph(slide, nodes, [("play", "movie", "rdf:type", "inferred")])
    P.line(slide, 7.1, 3.35, 7.9, 3.35, color=T.AMBER, width=2.5)
    P.text(slide, 8.15, 4.75, 3.8, 0.4, "a 1947 stage play, now typed as a movie", size=13,
           italic=True, color=T.CLAY, align=PP_ALIGN.CENTER)
    w = (CW - 0.3) / 2
    P.card(slide, L, 5.25, w, 1.4, "Why", "RDFS describes sets and how data flows between them. "
           "It is not validation.", fill=T.CLAY_TINT, title_size=16)
    P.card(slide, L + w + 0.3, 5.25, w, 1.4, "Do", "Use domain and range deliberately, as tools "
           "for classification and discovery.", title_size=16)
    P.notes(slide, """Domain and range are the hardest terms for newcomers because they look
        like object-oriented declarations. They never reject data: a triple that violates them
        just classifies its subject or object. Merging a theatre archive that reuses
        directedBy makes A Streetcar Named Desire a Movie. This is the real owlrl output.""")
    return slide


def s_rdfs_plus(prs, ex):
    slide = content(prs, "Part 5 · RDFS-Plus", "RDFS-Plus: let the engine work out sameness")
    sa = ex["same_as"]
    P.code(slide, L, 1.8, 6.3, 2.15, sa.code, label="Turtle", size=12)
    nodes = {
        "fdean": Node(8.35, 2.3, "fan:Dean_James", w=2.1),
        "dean": Node(11.85, 2.3, ":JamesDean", w=1.75),
        "pid": Node(10.1, 3.75, '"P-0015"', kind="literal", w=1.25),
    }
    P.graph(slide, nodes, [
        ("fdean", "pid", "fan:registryId", "asserted", (-0.45, 0)),
        ("dean", "pid", ":registryId", "asserted", (0.4, 0)),
        ("fdean", "dean", "owl:sameAs", "inferred"),
    ])
    P.text(slide, 7.25, 4.1, R - 7.25, 0.35, "one registry id → one individual", size=13,
           italic=True, color=T.AMBER, align=PP_ALIGN.CENTER)
    P.code(slide, L, 4.2, 6.3, 2.2, sa.facts["query"], label="SPARQL", size=12)
    x0 = 7.25
    P.text(slide, x0, 4.75, 2.2, 0.3, "Before alignment", size=13, bold=True, color=T.MUTED)
    P.text(slide, x0, 5.1, 2.2, 0.8, f"{sa.facts['before_rows']} rows", size=36, bold=True,
           color=T.CLAY, font=T.HEAD_FONT)
    P.line(slide, x0 + 2.0, 5.5, x0 + 2.5, 5.5, color=T.AMBER, width=2.5)
    x1 = x0 + 2.7
    P.text(slide, x1, 4.75, R - x1, 0.3, "After OWL RL inference", size=13, bold=True, color=T.MUTED)
    P.table(slide, x1, 5.15, R - x1, sa.header, sa.rows, col_widths=[0.95, R - x1 - 0.95], size=13,
            row_h=0.4)
    P.notes(slide, """RDFS-Plus adds sameness. equivalentClass and equivalentProperty align
        vocabularies; sameAs says two individuals are one. Declaring registryId inverse
        functional lets the engine compute sameness instead of requiring it to be asserted:
        a shared registry id makes the fan wiki's James Dean and the studio's the same
        person, so a question that spans both sources now returns an answer.""")
    return slide


# --------------------------------------------------------------------------- #
# Part 6: deep dive, vocabularies
# --------------------------------------------------------------------------- #
def s_wild(prs, ex):
    slide = content(prs, "Part 6 · RDFS-Plus in the wild",
                    "Real vocabularies apply RDFS-Plus selectively")
    w = (CW - 0.6) / 3
    cards = [
        ("Data.gov", "D", [
            "Hundreds of government datasets as raw RDF, prepared for no particular app.",
            "Use: mashups across datasets and the wider Web.",
            [("Lesson: ", {"bold": True, "color": T.TEAL}),
             ("SPARQL does most of the work, with patterns, aggregates, filters and unions.", {})]]),
        ("FOAF", "F", [
            "People, groups and documents, with one person's data spread across sites.",
            "Use: own and host your data, alongside social networks.",
            [("Lesson: ", {"bold": True, "color": T.TEAL}),
             ("Start with intuitive terms, then formalize them with RDFS-Plus over time.", {})]]),
        ("Open Graph Protocol", "O", [
            "About two dozen types and properties in one namespace, embedded with simplified RDFa.",
            "Use: “like” any page on the Web.",
            [("Lesson: ", {"bold": True, "color": T.TEAL}),
             ("Linking with rdfs:seeAlso commits to little, but is enough to avoid a silo.", {})]]),
    ]
    for i, (head, badge, body) in enumerate(cards):
        P.card(slide, L + i * (w + 0.3), 1.85, w, 3.15, head, body, badge=badge)
    P.text(slide, L, 5.25, 3.0, 0.45, "How strongly to link terms", size=14, bold=True,
           anchor=MSO_ANCHOR.MIDDLE)
    links = [("rdfs:seeAlso", "weak: “related”, as OGP does"),
             ("rdfs:subPropertyOf", "one-way: uses flow upward"),
             ("owl:equivalentProperty", "strong: interchangeable")]
    lx, lw = L + 3.0, (CW - 3.0 - 0.4) / 3
    for i, (term, desc) in enumerate(links):
        box = P.rect(slide, lx + i * (lw + 0.2), 5.12, lw, 0.76, fill=[T.WHITE, T.MIST, T.MINT_SOFT][i],
                     line=T.LINE, radius=0.1)
        P.shape_text(box, [[(term, {"bold": True, "font": T.CODE_FONT, "size": 13})],
                           [(desc, {"size": 13, "color": T.MUTED})]], align=PP_ALIGN.LEFT, margin=0.15)
    P.takeaway(slide, L, 6.15, CW,
               "Nonunique naming in practice: OGP reinvented FOAF terms and coexists instead of competing.")
    P.notes(slide, """Data.gov: raw government data whose RDF form enables mashups, with SPARQL
        doing most of the work. FOAF: people own their data; some terms are defined only in
        natural language, others formally with RDFS-Plus, tightened as they are used. Facebook's
        OGP kept one namespace for its content managers and links to other vocabularies with
        rdfs:seeAlso rather than equivalentProperty.""")
    return slide


def s_skos(prs, ex):
    slide = content(prs, "Part 6 · SKOS", "SKOS: vocabularies as linked data")
    nodes = {
        "libd": Node(1.75, 2.1, "lib:DramaFilms", kind="ghost", w=2.0, h=0.45),
        "libf": Node(5.85, 2.1, "lib:FrontierFilms", kind="ghost", w=2.25, h=0.45),
        "drama": Node(2.75, 3.6, "g:Drama", w=1.4),
        "western": Node(5.85, 3.6, "g:Western", w=1.5),
        "family": Node(1.6, 5.3, "g:FamilyDrama", w=1.85),
        "teen": Node(3.85, 5.3, "g:TeenDrama", w=1.7),
        "epic": Node(5.95, 5.3, "g:EpicWestern", w=1.85),
    }
    P.graph(slide, nodes, [
        ("drama", "libd", "exactMatch", "asserted", (0.2, 0)),
        ("western", "libf", "closeMatch"),
        ("family", "drama", "broader", "asserted", (-0.3, 0)),
        ("teen", "drama", "broader", "asserted", (0.3, 0)),
        ("epic", "western", "broader"),
    ], label_size=11)
    P.text(slide, L, 6.0, 6.5, 0.7,
           "Matching links two vocabularies while each term stays distinct in its own source.",
           size=14, color=T.MUTED)
    x0 = 7.4
    q = ex["skos"]
    P.code(slide, x0, 1.8, R - x0, 2.55, q.code, label="SPARQL", size=11)
    P.text(slide, x0, 4.5, R - x0, 0.3, "Which movies are dramas, at any depth?", size=13, bold=True,
           color=T.MUTED)
    P.table(slide, x0, 4.85, R - x0, q.header, q.rows, col_widths=[2.9, R - x0 - 2.9], size=13,
            row_h=0.36)
    P.text(slide, x0, 6.35, R - x0, 0.4, "skos:broader isn't transitive by default; the query opts in with *.",
           size=12, color=T.MUTED)
    P.notes(slide, """SKOS is the W3C Recommendation for publishing controlled vocabularies.
        Every concept gets a URI, so terms from different vocabularies can share a statement;
        exactMatch, closeMatch, broadMatch and narrowMatch relate them across sources. broader
        and narrower resemble subClassOf, but you choose whether to treat them as transitive.
        Here the property path broader* finds all three dramas.""")
    return slide


# --------------------------------------------------------------------------- #
# Part 7: deep dive, OWL
# --------------------------------------------------------------------------- #
def s_restrictions(prs, ex):
    slide = content(prs, "Part 7 · Basic OWL", "Restrictions define classes from what's known")
    hv = ex["has_value"]
    P.code(slide, L, 1.8, 6.4, 1.6, hv.code, label="Turtle")
    P.text(slide, L, 3.7, 6.4, 0.3, "Inferred by OWL RL", size=13, bold=True, color=T.MUTED)
    P.table(slide, L, 4.05, 6.4, hv.header, hv.rows, col_widths=[2.5, 1.4, 2.5], size=13, row_h=0.38,
            mono_cols=(0, 1, 2), header_fill=T.AMBER)
    P.text(slide, L, 5.85, 6.4, 0.8,
           "Membership is computed: any movie starring :JamesDean, from any source, joins the class.",
           size=14)
    x0 = 7.4
    kinds = [
        ("owl:hasValue", "Members have a particular value: starring :JamesDean."),
        ("owl:someValuesFrom", "At least one value comes from a class: starring some :Actor."),
        ("owl:allValuesFrom", "Every value, if any, comes from a class: genre only drama concepts."),
    ]
    for i, (head, body) in enumerate(kinds):
        P.card(slide, x0, 1.8 + i * 1.45, R - x0, 1.3, head, body, title_size=15)
    P.takeaway(slide, x0, 6.25, R - x0, "A restriction is a class, so others can reuse it.", size=14)
    P.notes(slide, """owl:Restriction defines a class by describing its members in terms of
        classes and properties already modeled, something RDFS cannot do. Because a restriction
        is itself a class, other sources can build on it. Here JamesDeanMovie is defined as
        everything starring James Dean, and OWL RL classifies all three films.""")
    return slide


def s_counting(prs, ex):
    slide = content(prs, "Part 7 · Counting and sets",
                    "Counting needs a closed world: close it on purpose")
    tiles = [
        ("unionOf", "either class"), ("intersectionOf", "both classes"),
        ("complementOf", "not in the class"), ("cardinality", "min, max, exact"),
        ("disjointWith", "no shared members"), ("differentFrom", "these names differ"),
    ]
    tw = (5.6 - 0.15) / 2
    for i, (name, desc) in enumerate(tiles):
        x = L + (i % 2) * (tw + 0.15)
        y = 1.8 + (i // 2) * 1.05
        box = P.rect(slide, x, y, tw, 0.9, fill=T.MIST, radius=0.1)
        P.shape_text(box, [[(f"owl:{name}", {"bold": True, "font": T.CODE_FONT, "size": 13})],
                           [(desc, {"color": T.MUTED, "size": 13})]], align=PP_ALIGN.LEFT, margin=0.18)
    P.text(slide, L, 5.1, 5.6, 1.6,
           "Open world and nonunique names make OWL cautious: two names may be one thing, and "
           "more facts may turn up.", size=15)
    mc = ex["max_card"]
    x0 = 6.7
    P.code(slide, x0, 1.8, R - x0, 2.0, mc.code, label="Turtle", size=12)
    P.text(slide, x0, 4.0, R - x0, 0.3, "Two directors for Giant? OWL concludes:", size=13, bold=True,
           color=T.MUTED)
    P.pill(slide, x0, 4.35, R - x0, 0.48, ":GeorgeStevens owl:sameAs fan:G_Stevens", T.AMBER_TINT,
           color=T.INK, line=T.AMBER, size=13, font=T.CODE_FONT)
    P.text(slide, x0, 5.1, R - x0, 0.3, "Close the world: declare the names different", size=13,
           bold=True, color=T.MUTED)
    P.pill(slide, x0, 5.45, R - x0, 0.75,
           "OWL RL reports an inconsistency: sameAs and differentFrom on the same pair",
           T.CLAY_TINT, color=T.INK, line=T.CLAY, size=14)
    P.notes(slide, """OWL adds unions, intersections, complements, cardinality and disjointness,
        so a model can rule things out. Counting is hard on the Web: a max-cardinality of one
        with two named directors doesn't produce an error, it makes the names the same person.
        Only once we close the world, by declaring them different, does the engine report a
        contradiction. Both results come from owlrl.""")
    return slide


def s_ontologies(prs, ex):
    slide = content(prs, "Part 7 · Ontologies on the Web",
                    "Ontologies on the Web span a data spectrum")
    P.line(slide, L + 0.3, 2.15, R - 0.3, 2.15, color=T.TEAL, width=2.5)
    P.text(slide, L, 1.7, 5, 0.35, "almost no data in the ontology", size=13, color=T.MUTED)
    P.text(slide, R - 5, 1.7, 5, 0.35, "very large, richly connected data", size=13, color=T.MUTED,
           align=PP_ALIGN.RIGHT)
    w = (CW - 0.6) / 3
    cards = [
        ("Good Relations", [
            "Small vocabulary for business offerings, in RDFa on product pages",
            "Links to UN/CEFACT unit codes",
            "Used on tens of thousands of sites; richer search results",
        ]),
        ("QUDT (NASA)", [
            "A global URI per unit, so “pound” is unambiguous",
            "Converts between commensurate units",
            "Dimensional analysis catches kilograms + centimetres",
        ]),
        ("OBO ontologies", [
            "Unambiguous names for tens of thousands of biological concepts (e.g. ChEBI)",
            "Organize lab data and locate publications",
            "A Rosetta Stone across naming schemes",
        ]),
    ]
    for i, (head, body) in enumerate(cards):
        x = L + i * (w + 0.3)
        P.circle(slide, x + w / 2, 2.15, 0.34, T.AMBER)
        P.card(slide, x, 2.6, w, 3.0, head, body, bullets=True)
    P.takeaway(slide, L, 5.95, CW, "Shared role: give standard entities globally unambiguous names.")
    P.notes(slide, """Good Relations has almost no data of its own: the data is spread across the
        Web in RDFa. QUDT sits in the middle, holding some data but mainly connecting data sets
        through units, conversions and dimensions. The OBO ontologies are very large and richly
        connected, a Rosetta Stone for biological names.""")
    return slide


# --------------------------------------------------------------------------- #
# Part 8: deep dive, modeling well
# --------------------------------------------------------------------------- #
def s_practice(prs, ex):
    slide = content(prs, "Part 8 · Good and bad practice",
                    "Good models start from competency questions")
    rules = [
        ("Know what you want", "Competency questions set the scope; expect unanticipated merges too."),
        ("Say what you mean", "Avoid wishful naming: check that inferences match intent."),
        ("Model for reuse", "Clear names plus rdfs:label, rdfs:comment, rdfs:seeAlso."),
        ("Classes vs individuals", "Can't imagine instances? Not a class. Know its values? An instance."),
    ]
    for i, (head, body) in enumerate(rules):
        y = 1.8 + i * 1.2
        P.circle(slide, L + 0.25, y + 0.25, 0.5, T.TEAL, str(i + 1), size=16)
        P.text(slide, L + 0.7, y + 0.03, 4.6, 0.35, head, size=16, bold=True)
        P.text(slide, L + 0.7, y + 0.42, 4.6, 0.75, body, size=14, color=T.MUTED)
    x0 = 6.0
    cq = ex["cq"]
    P.text(slide, x0, 1.8, R - x0, 0.35, "Competency questions as executable tests", size=16, bold=True)
    colors = {(r, 2): (T.TEAL_DARK if row[2] == "pass" else T.CLAY) for r, row in enumerate(cq.rows)}
    fills = {(r, 2): (T.MINT_SOFT if row[2] == "pass" else T.CLAY_TINT) for r, row in enumerate(cq.rows)}
    P.table(slide, x0, 2.3, R - x0, cq.header, cq.rows, col_widths=[2.6, 3.4, R - x0 - 6.0], size=13,
            row_h=0.55, mono_cols=(1,), cell_colors=colors, cell_fills=fills)
    P.text(slide, x0, 5.25, R - x0, 1.4,
           "The failing question is a finding, not a bug: the model has no awards yet. Testing also "
           "means checking how inferences behave on data from other sources.", size=14)
    P.notes(slide, """Start from competency questions, the questions the model must answer, and
        model no further than needed, while remembering a semantic model will be merged with
        unanticipated sources. The table runs four questions as SPARQL ASK queries against the
        inferred studio graph; three pass and the awards question exposes a scope gap.""")
    return slide


def s_antipatterns(prs, ex):
    slide = content(prs, "Part 8 · Good and bad practice", "Four antipatterns to avoid")
    w, h = (CW - 0.3) / 2, 2.3
    cards = [
        ("Rampant classism", "Making everything a class, driven by tooling or thesaurus conversion.",
         "e.g. a class :JamesDean, though he has a registry id: an instance."),
        ("Exclusivity", "Assuming only known members of a superclass can belong to a subclass.",
         "e.g. treating the three listed films as the only possible :JamesDeanMovie."),
        ("Objectification", "Treating classes as templates. Properties stand alone; classes are sets.",
         "e.g. thinking :starring belongs to :Movie, so nothing else may use it."),
        ("Creeping conceptualization", "“Design for reuse” is not “say everything you can”.",
         "Warning sign: :FeatureFilm, :ColorFeatureFilm, … with no instances."),
    ]
    for i, (head, body, eg) in enumerate(cards):
        x = L + (i % 2) * (w + 0.3)
        y = 1.8 + (i // 2) * (h + 0.25)
        P.card(slide, x, y, w, h, head, [body, [(eg, {"italic": True, "color": T.MUTED})]],
               badge=i + 1, badge_fill=T.CLAY, body_size=15)
    P.notes(slide, """None of these is strictly an error under AAA, but each defeats sharing.
        Rampant classism breaks under inference. Exclusivity is consistent with the Semantic Web
        assumptions and fixable with a simple pattern. Objectification clashes with them, since
        properties exist independently of classes. Creeping conceptualization restricts others;
        keep it simple.""")
    return slide


def s_profiles(prs, ex):
    slide = content(prs, "Part 8 · Expert modeling in OWL",
                    "OWL 2 profiles: provable or practical")
    P.card(slide, L, 1.8, 5.2, 2.25, "Provable",
           "A computer can always decide, in finite time, whether two sources mean the same thing. "
           "That needs carefully restricted logic.", badge="1")
    P.card(slide, L, 4.3, 5.2, 2.25, "Practical",
           "Treat the model like a program: test and fix until it behaves on the data that matters. "
           "Good answers most of the time.", badge="2", badge_fill=T.AMBER)
    x0 = 6.2
    P.text(slide, x0, 1.8, R - x0, 0.4, "Same vocabulary everywhere; only the usage rules differ",
           size=15, bold=True)
    rows = [
        ("OWL 2 DL", "Most expressive while keeping provable answers"),
        ("OWL 2 QL", "Applications on top of ordinary databases"),
        ("OWL 2 EL", "Very large ontologies, as in the life sciences"),
        ("OWL 2 RL", "Rule-based software (owlrl, which built this deck)"),
    ]
    P.table(slide, x0, 2.35, R - x0, ("Profile", "Built for"), rows, col_widths=[1.6, R - x0 - 1.6],
            size=15, row_h=0.62, cell_colors={(r, 0): T.TEAL_DARK for r in range(len(rows))})
    P.takeaway(slide, x0, 5.85, R - x0, "Choose by what you need: proof, or efficient integration.",
               size=14)
    P.notes(slide, """OWL is one toolbox; each subset has the same tools with different usage
        rules. Some people want provable models; others treat models as programs that merge and
        reshape data, accepting good answers most of the time. DL keeps guarantees, QL sits on
        databases, EL scales to huge ontologies, RL suits rule engines like owlrl.""")
    return slide


def s_owl2(prs, ex):
    slide = content(prs, "Part 8 · Expert modeling in OWL",
                    "New in OWL 2: chains, keys and punning, plus rules")
    ch = ex["chain"]
    P.code(slide, L, 1.8, 6.4, 0.95, ch.code, label="Turtle", size=12)
    nodes = {
        "dean": Node(1.6, 4.85, ":JamesDean", w=1.8),
        "giant": Node(3.8, 3.45, ":Giant", w=1.25),
        "gs": Node(6.0, 4.85, ":GeorgeStevens", w=1.95),
    }
    P.graph(slide, nodes, [
        ("dean", "giant", ":starredIn", "asserted", (-0.35, -0.1)),
        ("giant", "gs", ":directedBy", "asserted", (0.35, -0.1)),
        ("dean", "gs", ":workedWith", "inferred"),
    ])
    P.text(slide, L, 5.5, 6.4, 1.1, [[
        (f"{len(ch.rows)} :workedWith triples for James Dean, ", {}),
        ("identical to the CONSTRUCT rule in Part 4", {"bold": True}),
        (" (checked at build time).", {})]], size=14)
    x0 = 7.4
    feats = [
        ("Keys (owl:hasKey)", "Same name and address ⇒ same person, like a multi-column key. "
                              "Still “same”, never “probably same”."),
        ("Punning", "One name as class and individual. Safer: keep them separate and relate them."),
        ("Rules and RIF", "IF/THEN rules fill OWL's gaps, but big rule bases are hard to maintain. "
                          "OWL 2 RL and RIF combine both."),
    ]
    for i, (head, body) in enumerate(feats):
        P.card(slide, x0, 1.8 + i * 1.62, R - x0, 1.47, head, body, title_size=16)
    P.notes(slide, """OWL 2 adds property chains (uncle is brother of a parent; here workedWith is
        starredIn then directedBy), keys built from several properties, and punning, which lets
        one name be a class and an individual. Rules fill remaining gaps but are hard to maintain
        at scale and carry no guarantees; OWL 2 RL and RIF combine the two.""")
    return slide


# --------------------------------------------------------------------------- #
# Part 9: conclusions
# --------------------------------------------------------------------------- #
def s_conclusions(prs, ex):
    slide = content(prs, "Part 9 · Conclusions", "What's new isn't the logic, it's the Web")
    P.text(slide, L, 1.7, CW, 0.5,
           "“Same old, same old”? Classes, subclasses and properties are decades old. The Semantic "
           "Web is part of the Web's infrastructure, not an app running on it.",
           size=16, italic=True, color=T.MUTED)
    w = (CW - 0.6) / 3
    cards = [
        ("Global references", "URIs let separately built models merge by mentioning each "
                              "other's resources in one statement."),
        ("Openness", "The Web tolerates broken or wrong content, so inference stays "
                     "conservative under the open world assumption."),
        ("Scale", "No store holds the whole Web, and none needs to: an app merges any page, "
                  "not every page."),
    ]
    for i, (head, body) in enumerate(cards):
        P.card(slide, L + i * (w + 0.3), 2.55, w, 2.45, head, body, badge=i + 1, body_size=16)
    P.takeaway(slide, L, 5.45, CW,
               "The craft: combine RDF, RDFS and OWL into a system that data can flow through.")
    P.notes(slide, """Veterans call it same old, same old, but the novelty is the Web. Like the
        Web, the Semantic Web wins by being open: global URIs let models merge, openness forces
        conservative open-world inference, and applications only need to merge the pages they
        use.""")
    return slide


def s_takeaways(prs, ex):
    slide = content(prs, "Closing notes", "What I'd do with this as a practitioner")
    steps = [
        ("Model with RDF, query with SPARQL", "URIs, triples and graph patterns can represent and query a curated domain."),
        ("Add inference where needed", "subClassOf, subPropertyOf and sameAs do most of the integration work."),
        ("Test with competency questions", "ASK queries run on every build keep the model honest."),
        ("Be careful with automatic discovery", "OWL derives new facts, but automatic knowledge discovery isn't ready."),
    ]
    w = (CW - 0.9) / 4
    P.line(slide, L + 0.3, 2.15, R - w + 0.3, 2.15, color=T.LINE, width=2, arrow=False)
    for i, (head, body) in enumerate(steps):
        x = L + i * (w + 0.3)
        P.circle(slide, x + 0.3, 2.15, 0.6, T.TEAL, str(i + 1), size=18)
        P.text(slide, x, 2.7, w, 0.8, head, size=17, bold=True)
        P.text(slide, x, 3.5, w, 1.5, body, size=15, color=T.MUTED)
    box = P.rect(slide, L, 4.9, CW, 1.2, fill=T.AMBER_TINT, radius=0.14)
    P.shape_text(box, [[("For an LLM application: ", {"bold": True}),
                        ("the LLM brings the reasoning; the ontology's job is a well-curated, linked "
                         "dataset with clear, global names.", {})]],
                 size=16, align=PP_ALIGN.LEFT, margin=0.3)
    P.notes(slide, """My reading: build a bounded domain model in RDF and query it with SPARQL;
        add light inference where integration needs it; turn competency questions into tests;
        and treat automatic knowledge discovery with OWL as not yet ready.""")
    return slide


def s_references(prs, ex):
    slide = content(prs, "References", "Sources and further reading")
    refs = [
        ("Main reference", [
            ("Allemang, D. and Hendler, J. (2011). ", {}),
            ("Semantic Web for the Working Ontologist: Effective Modeling in RDFS and OWL",
             {"italic": True}),
            (", 2nd ed. Morgan Kaufmann.", {})]),
        ("Definition of an ontology", [
            ("Gruber, T. (1993). A translation approach to portable ontology specifications. ", {}),
            ("Knowledge Acquisition", {"italic": True}), (" 5(2). Studer, R., Benjamins, V. R. and "
             "Fensel, D. (1998). Knowledge engineering: principles and methods. ", {}),
            ("Data & Knowledge Engineering", {"italic": True}), (" 25(1–2).", {})]),
        ("Semantic Web Stack image", [
            ("W3C, public domain, via Wikimedia Commons: "
             "commons.wikimedia.org/wiki/File:Semantic_Web_Stack.png.", {})]),
        ("W3C standards", [
            ("RDF 1.1 Concepts, RDF Schema 1.1, SPARQL 1.1 Query, SKOS Reference, OWL 2 Overview, "
             "RIF Overview, SHACL, PROV-O and RDF Dataset Canonicalization, all at w3.org/TR.", {})]),
        ("Software", [
            ("rdflib, and owlrl (an OWL 2 RL reasoner), run every example in this deck.", {})]),
    ]
    for i, (head, body) in enumerate(refs):
        y = 1.8 + i * 1.0
        P.circle(slide, L + 0.25, y + 0.25, 0.5, T.TEAL, str(i + 1), size=16)
        P.text(slide, L + 0.7, y + 0.03, CW - 0.7, 0.35, head, size=16, bold=True)
        P.text(slide, L + 0.7, y + 0.42, CW - 0.7, 0.7, [body], size=14, color=T.MUTED)
    P.notes(slide, """The main reference is Semantic Web for the Working Ontologist, second
        edition. The definition of an ontology is Gruber's, as refined by Studer, Benjamins and
        Fensel. The layer-cake diagram is the W3C Semantic Web Stack from Wikimedia Commons.
        The standards themselves are W3C Recommendations.""")
    return slide


GLOSSARY = [
    ("Ontology", "A formal, explicit, shared model of a domain: classes, properties, axioms."),
    ("Class", "A kind of thing, understood as the set of its members."),
    ("Individual", "A particular thing; an instance of one or more classes."),
    ("Property", "A relationship between individuals, or from an individual to a value."),
    ("Axiom", "A statement about classes or properties that a reasoner can use."),
    ("Triple", "One RDF statement: subject, predicate, object."),
    ("URI / IRI", "A global name for a thing, class or property."),
    ("Literal", "A plain value, such as a string or number, with optional datatype."),
    ("TBox / ABox", "The model (terminology) and the data (assertions)."),
    ("Knowledge graph", "An ontology plus the data described with it, in one graph."),
    ("Inference", "Deriving triples that follow from the ones stated."),
    ("Reasoner", "Software that performs inference, e.g. owlrl."),
    ("Asserted / inferred", "Stated by a source / derived by inference. Same status once present."),
    ("Open world", "A missing fact is unknown, not false."),
    ("Nonunique naming", "One thing may have several names; never assume two names are two things."),
    ("Vocabulary", "A set of terms, usually under one namespace, such as FOAF or SKOS."),
    ("Taxonomy", "Terms arranged in a broader–narrower hierarchy."),
    ("Profile", "A subset of OWL 2 (DL, EL, QL, RL) with known reasoning cost."),
    ("Competency question", "A question the model must be able to answer; a test."),
    ("SPARQL endpoint", "A web service that answers SPARQL queries over a store."),
]


def s_glossary(prs, ex):
    slide = content(prs, "Appendix", "Glossary")
    half = (CW - 0.3) / 2
    n = len(GLOSSARY) // 2
    for c in range(2):
        rows = GLOSSARY[c * n:(c + 1) * n]
        P.table(slide, L + c * (half + 0.3), 1.7, half, ("Term", "Meaning"), rows,
                col_widths=[1.9, half - 1.9], size=12, row_h=0.47,
                cell_fills={(r, 0): T.MIST for r in range(len(rows))})
    P.notes(slide, """Terms used throughout the deck, in roughly the order a newcomer meets
        them.""")
    return slide



# --------------------------------------------------------------------------- #
# deck order
# --------------------------------------------------------------------------- #
def part(index):
    return lambda prs, ex: divider(prs, index)


DECK = [
    s_title,
    s_definition, s_anatomy, s_spectrum, s_vs_database,
    part(1), s_why, s_aaa, s_modeling,
    part(2), s_stack, s_stack_w3c, s_beyond_stack,
    s_layer_uri, s_layer_rdf, s_layer_rdfs, s_layer_rdfs_plus, s_layer_owl, s_layer_skos,
    s_layer_rules, s_layer_sparql,
    s_together_flow, s_together_query, s_cheatsheet,
    part(3), s_table_to_triples, s_merge, s_identity, s_architecture, s_sparql_select, s_sparql_explore,
    part(4), s_inference, s_when, s_rdfs, s_domain_trap, s_rdfs_plus,
    part(5), s_wild, s_skos,
    part(6), s_restrictions, s_counting, s_ontologies,
    part(7), s_practice, s_antipatterns, s_profiles, s_owl2,
    part(8), s_conclusions, s_takeaways, s_references, s_glossary,
]
